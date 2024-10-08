import pandas as pd
import requests
import base64
import json
import psycopg2 as ps
import csv
import time
import os
import concurrent.futures
from datetime import datetime
from openpyxl import load_workbook

from CSV_Files_Path import base_dir, login_csv, basic_user_info, refund_demand_path, tax_deposit_path, recent_filed_return
from CSV_Files_Path import recent_form_filed_path, active_bank_path, inactive_bank_path, failed_bank_path, juris_diction
from CSV_Files_Path import src_business_salaried, src_house_property, authorised_signature, representative_asseses, active_demat_account
from CSV_Files_Path import fya_notice_count, fya_notice_description, fya_all_notices_path, fya_notices_letter, fya_download_dir
from CSV_Files_Path import fyi_notice_count, fyi_notice_description, fyi_all_notices, fyi_notices_letter, fyi_download_dir

from Error_File_Path import Error_process_login, Error_user_info, Error_refund_demand, Error_tax_deposit, Error_recent_file_return
from Error_File_Path import Error_recent_form_filed, Error_active_bank, Error_inactive_bank, Error_failed_bank
from Error_File_Path import Error_juris_diction, Error_src_business_salaried, Error_src_house_property
from Error_File_Path import Error_authorised_signature, Error_representative_asses, Error_active_demat_account, Error_demat_account
from Error_File_Path import Error_fya, Error_fya_notice_count, Error_fya_notice_description, Error_fya_all_notices, Error_fya_notice_letter, Error_fya_notice_download
from Error_File_Path import Error_fyi, Error_fyi_count, Error_fyi_notice_description, Error_fyi_all_notices, Error_fyi_notice_letter, Error_fyi_notice_download, Error_direct_notice_download_path

#----------------------------------------All Links----------------------------------------
start = time.perf_counter()

loginid_url = "https://eportal.incometax.gov.in/iec/loginapi/login"
loginpassword_url = "https://eportal.incometax.gov.in/iec/loginapi/login"

userinfo_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/saveEntity"
tax_deposit_url = "https://eportal.incometax.gov.in/iec/efileprocessingapi/auth/saveEntity"
recent_filed_return_url = "https://eportal.incometax.gov.in/iec/itrweb/auth/v0.1/returns/previous"
recent_forms_filed_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/saveEntity"
rep_assess_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
refund_demand_url = "https://eportal.incometax.gov.in/iec/itrweb/auth/v0.1/returns/demand/refund"
refund_demand_info = "https://eportal.incometax.gov.in/iec/loginapi/auth/saveEntity" 
income_source_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
jurisdiction_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/saveEntity"
inactive_bank_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
failed_bank_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
demat_acc_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
auth_sign_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
active_demat_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"
bank_info_url = "https://eportal.incometax.gov.in/iec/servicesapi/auth/getEntity"

fya_count_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fya_notice_description_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fya_all_notices_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fya_notice_letter_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/saveEntity"
fya_download_notice_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/saveEntity" 

fyi_notice_count_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fyi_notice_description_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fyi_all_notices_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"
fyi_notice_letter_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/saveEntity"
fyi_download_notice_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/saveEntity"
fyi_direct_notice_download_url = "https://eportal.incometax.gov.in/iec/returnservicesapi/auth/getEntity"

logout_url = "https://eportal.incometax.gov.in/iec/loginapi/login"

ca_registration_number='155154W'

current_year = datetime.now().year
last_year = current_year - 1

#---------------------------------------- DATABASE CONNECTION----------------------------------------
host_name_ef_database = 'itr-db-2.ch8ys6oq2fb7.ap-south-1.rds.amazonaws.com'
dbname_ef_database = 'itr_database_2'
port_ef_database = '5432'
username_ef_database = 'postgres'
password_ef_database = 'postgres'

def connect_to_ef_database(host_name_ef_database, dbname_ef_database, port_ef_database, username_ef_database, password_ef_database):
    try:
        conn_ef_db = ps.connect(host=host_name_ef_database, database=dbname_ef_database, port=port_ef_database, user=username_ef_database, password=password_ef_database)
    
    except ps.OperationalError as e:
        raise e
    
    else:
        return conn_ef_db

conn_ef_db = connect_to_ef_database(host_name_ef_database, dbname_ef_database, port_ef_database, username_ef_database, password_ef_database)
curr_ef_db = conn_ef_db.cursor()

#----------------------------------------DEFINE FUNCTION----------------------------------------
def process_login(id_password):
    username , password = id_password
    try:
        time.sleep(3)
        #--------------------PASS THE ID--------------------
        payload_loginId = json.dumps({
        "entity": username,
        "serviceName": "wLoginService"
        })
        
        headers_loginId = {
        'Accept': 'application/json, text/plain, */*',
        'Content-Type': 'application/json',
        'sn': 'wLoginService',
        }

        loginId = requests.request("POST", loginid_url, headers = headers_loginId, data = payload_loginId)
        time.sleep(5)
    
        #--------------------GET THE REQUEST ID AND PAN NUMBER--------------------
        response_loginId = json.loads(loginId.text)
        response_Id = response_loginId.get('reqId')
        response_entity = response_loginId.get('entity')
        response_role = response_loginId.get('role')

        #--------------------ENCODE THE PASSWORD--------------------
        passwords_bytes = password.encode("ascii")
        password_base_64 = base64.b64encode(passwords_bytes)
        password_base_64_string = password_base_64.decode("ascii")
        time.sleep(5)

        #--------------------PASS THE ENCODED PASSWORD, ID AND PAN NUMBER--------------------
        payload_loginPassword = json.dumps({
        "errors": [],
        "reqId": response_Id,
        "entity": response_entity,
        "entityType": "PAN",
        "role": response_role,
        "uidValdtnFlg": "true",
        "aadhaarMobileValidated": "false",
        "secAccssMsg": "",
        "secLoginOptions": "",
        "dtoService": "LOGIN",
        "exemptedPan": "false",
        "userConsent": "N",
        "imgByte": None,
        "pass": password_base_64_string,
        "passValdtnFlg": None,
        "otpGenerationFlag": None,
        "otp": None,
        "otpValdtnFlg": None,
        "otpSourceFlag": None,
        "contactPan": None,
        "contactMobile": None,
        "contactEmail": None,
        "email": None,
        "mobileNo": None,
        "forgnDirEmailId": None,
        "imagePath": None,
        "serviceName": "loginService"
        })

        headers_loginPassword = {
        'Accept': 'application/json, text/plain, */*',
        'Accept-Language': 'en-US,en;q=0.9',
        'Connection': 'keep-alive',
        'Content-Type': 'application/json',
        'Cookie': 'rxVisitor=1709531052663DL75B9Q4S1IMFJFGDA9346K3RF5N9SID; rxvt=1709532854581|1709531052664; dtPC=-23$131052659_893h-vVMDPAFUDRRITORLOUAKLUHUKUSHRIRKA-0e0; dtCookie=v_4_srv_2_sn_DEBEG43JVFPPM5RM3NQTLG1PI0BCFKG0_perc_100000_ol_0_mul_1_app-3A789f7f41ec9fa652_0; dtSa=true%7CC%7C-1%7CLogin%7C-%7C1709531078781%7C131052659_893%7Chttps%3A%2F%2Fwww.incometax.gov.in%2Fiec%2Ffoportal%2F%7C%7C%7C%7C; 693c4e2771754eedb1d75ba0debd40d8=3d4aca42112fca0abacdec6024becf03; 4a75cee7266fb5ae654dc5e51e6a9fe3=8cebad70510e63e257eecb181e455b78; 735f08284897d8257a5156fc6c214f76=836e17132580e71d9d0420a39ebbf16b; 64912cb2ff2ddd44f1b8a0441cd026fb=2e567ea96c73979a5bf948d8a77e6b64; 2dcbb10b317da85bc2e359274540dc79=a416dbe57cf63de35afc0bf9b0f30aa3; cb75fe3af6a15223fa5633039e60ed6f=4935fd1049d10ea86f7c74991aa35697; eae3f7faaf0d5a512766019097a0d512=9eb16b73098a2aeb298f4297d07f7b5d; 1580efbee52cfd917d7413e4b29c4e2e=f2f31a1f76325305c3a109cf31df7e80; 49f6fb1f55554861258cb8d75043adee=3b7592f46113bf4bd6e386c423d38974; cda5f4d60b756653ef27a4733f4617c6=91c78cc681b7535c21b0c1eb479cfbd0; AuthToken=aad6c367124242aeb62b280cc3e3a5b4; 4a75cee7266fb5ae654dc5e51e6a9fe3=806c6b373b539eb401b1c9a744cccfdc',
        'Origin': 'https://eportal.incometax.gov.in',
        'Referer': 'https://eportal.incometax.gov.in/iec/foservices/',
        'Sec-Fetch-Dest': 'empty',
        'Sec-Fetch-Mode': 'cors',
        'Sec-Fetch-Site': 'same-origin',
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.0.0 Safari/537.36',
        'sec-ch-ua': '"Not A(Brand";v="99", "Google Chrome";v="121", "Chromium";v="121"',
        'sec-ch-ua-mobile': '?0',
        'sec-ch-ua-platform': '"Windows"',
        'sn': 'loginService'
        }
        
        login_Password = requests.request("POST", loginpassword_url, headers = headers_loginPassword, data = payload_loginPassword)
        time.sleep(5)
    
        #--------------------GET THE AUTH TOKEN--------------------
        if login_Password.status_code == 200:
            auth_token = login_Password.cookies.get('AuthToken')

        No_Belongs_to = {
            "2":"Spouse", "26":"Tax Return Preparer", "20":"Parent", "6":"Daughter", "24":"Authorized Representative", "8":"Sister",
            "7":"Brother", "25":"e-Return Intermediary", "22":"Friend", "1":"Self", "27":"Act as Authorised Representative",
            "5":"Son", "21":"Relative", "24": "Authorized Representative", "8": "Sister", "7": "Brother", "22": "Friend",
            "1": "Self","27": "Act as Authorised Representative","5": "Son","21": "Relative", "2": "Spouse","20": "Parent","6": "Daughter"
        }

        resdiantial_status = {
            "RES":"Resident", "NRI":"Non Resident"
        }

        country_code = {
            "5":"ITALY", "7":"KAZAKHSTAN", "8":"RUSSIAN FEDERATION", "9":"CHRISTMAS ISLAND", "14":"PORTUGAL", "15":"NORFOLK ISLAND",
            "20":"EGYPT", "28":"SOUTH AFRICA", "30":"GREECE", "31":"NETHERLANDS", "32":"BELGIUM", "33":"FRANCE", "35":"SPAIN", 
            "36":"HUNGARY",  "40":"ROMANIA",  "41":"SWITZERLAND",  "43":"AUSTRIA",  "44":"UNITED KINGDOM",  "45":"DENMARK", "46":"SWEDEN",
            "47":"NORWAY", "48":"POLAND", "49":"GERMANY", "51":"PERU", "52":"MEXICO", "53":"CUBA", "54":"ARGENTINA", "55":"BRAZIL",
            "56":"CHILE", "57":"COLOMBIA", "60":"MALAYSIA", "61":"AUSTRALIA", "62":"INDONESIA", "63":"PHILIPPINES", "64":"NEW ZEALAND",
            "65":"SINGAPORE", "66":"THAILAND", "81":"JAPAN", "82":"KOREA REPUBLIC OF", "86":"CHINA", "90":"TURKEY", "91":"INDIA",
            "92":"PAKISTAN", "93":"AFGHANISTAN", "94":"SRI LANKA", "95":"MYANMAR", "98":"IRAN ISLAMIC REPUBLIC OF", "211":"SOUTH SUDAN", 
            "212":"MOROCCO", "213":"ALGERIA", "216":"TUNISIA", "218":"LIBYA", "220":"GAMBIA", "221":"SENEGAL", "222":"MAURITANIA", 
            "223":"MALI", "224":"GUINEA", "225":"COTE D'IVOIRE", "226":"BURKINA FASO", "227":"NIGER", "228":"TOGO", "229":"BENIN", 
            "230":"MAURITIUS", "231":"LIBERIA", "232":"SIERRA LEONE", "233":"GHANA", "234":"NIGERIA", "235":"CHAD",
            "236":"CENTRAL AFRICAN REPUBLIC", "237":"CAMEROON", "238":"CAPE VERDE", "239":"SAO TOME AND PRINCIPE", "240":"EQUATORIAL GUINEA", 
            "241":"GABON", "242":"CONGO", "243":"CONGO THE DEMOCRATIC REPUBLIC OF THE", "244":"ANGOLA", "245":"GUINEA-BISSAU", "248":"SEYCHELLES",
            "249":"SUDAN", "250":"RWANDA", "251":"ETHIOPIA", "252":"SOMALIA", "253":"DJIBOUTI", "254":"KENYA", "255":"TANZANIA UNITED REPUBLIC OF",
            "256":"UGANDA", "257":"BURUNDI", "258":"MOZAMBIQUE", "261":"MADAGASCAR", "262":"REUNION", "264":"NAMIBIA", "265":"MALAWI",
            "266":"LESOTHO", "267":"BOTSWANA", "268":"SWAZILAND", "269":"MAYOTTE", "270":"COMOROS", "290":"SAINT HELENA ASCENSION AND TRISTAN DA CUNHA", 
            "291":"ERITREA", "297":"ARUBA", "298":"FAROE ISLANDS", "299":"GREENLAND", "350":"GIBRALTAR", "352":"LUXEMBOURG", "353":"IRELAND",
            "354":"ICELAND", "355":"ALBANIA", "356":"MALTA", "357":"CYPRUS", "358":"FINLAND", "359":"BULGARIA", "370":"LITHUANIA",
            "371":"LATVIA", "372":"ESTONIA", "1":"CANADA", "2":"UNITED STATES", "373":"MOLDOVA REPUBLIC OF", "374":"ARMENIA", "375":"BELARUS",
            "376":"ANDORRA", "377":"MONACO", "378":"SAN MARINO", "380":"UKRAINE", "381":"SERBIA", "382":"MONTENEGRO", "385":"CROATIA",
            "386":"SLOVENIA", "387":"BOSNIA AND HERZEGOVINA", "389":"MACEDONIA THE FORMER YUGOSLAV REPUBLIC OF", "420":"CZECH REPUBLIC", "421":"SLOVAKIA",
            "423":"LIECHTENSTEIN", "500":"FALKLAND ISLANDS (MALVINAS)", "501":"BELIZE", "502":"GUATEMALA", "503":"EL SALVADOR", "504":"HONDURAS",
            "505":"NICARAGUA", "506":"COSTA RICA", "507":"PANAMA", "508":"SAINT PIERRE AND MIQUELON", "509":"HAITI", "590":"GUADELOUPE",
            "591":"BOLIVIA (PLURINATIONAL STATE OF)", "592":"GUYANA","593":"ECUADOR","594":"FRENCH GUIANA","595":"PARAGUAY","596":"MARTINIQUE",
            "597":"SURINAME", "670":"TIMOR-LESTE", "672":"COCOS (KEELING) ISLANDS", "673":"BRUNEI DARUSSALAM", "674":"NAURU", "675":"PAPUA NEW GUINEA", 
            "676":"TONGA", "677":"SOLOMON ISLANDS", "679":"FIJI", "680":"PALAU", "682":"COOK ISLANDS", "683":"NIUE", "684":"AMERICAN SAMOA",
            "685":"SAMOA", "686":"KIRIBATI", "687":"NEW CALEDONIA", "688":"TUVALU", "689":"FRENCH POLYNESIA", "690":"TOKELAU", 
            "691":"MICRONESIA FEDERATED STATES OF","692":"MARSHALL ISLANDS","850":"KOREA DEMOCRATIC PEOPLE'S REPUBLIC OF", "852":"HONG KONG",
            "853":"MACAO", "855":"CAMBODIA", "856":"LAO PEOPLE 'S DEMOCRATIC REPUBLIC", "880":"BANGLADESH", "886":"TAIWAN",
            "960":"MALDIVES", "961":"LEBANON", "962":"JORDAN", "963":"SYRIAN ARAB REPUBLIC", "964":"IRAQ", "965":"KUWAIT", "966":"SAUDI ARABIA",
            "968":"OMAN", "970":"PALESTINE STATE OF", "971":"UNITED ARAB EMIRATES", "972":"ISRAEL", "973":"BAHRAIN", "974":"QATAR", "975":"BHUTAN",
            "976":"MONGOLIA", "977":"NEPAL", "992":"TAJIKISTAN", "993":"TURKMENISTAN", "994":"AZERBAIJAN", "995":"GEORGIA", "996":"KYRGYZSTAN",
            "1006":"SAINT BARTHELEMY", "1007":"SAINT MARTIN (FRENCH PART)", "1009":"UNITED STATES MINOR OUTLYING ISLANDS",
            "1012":"SVALBARD AND JAN MAYEN", "1014":"BRITISH INDIAN OCEAN TERRITORY", "1015":"CURACAO", "1242":"BAHAMAS",
            "1246":"BARBADOS", "1264":"ANGUILLA", "1268":"ANTIGUA AND BARBUDA", "1284":"VIRGIN ISLANDS BRITISH", "1340":"VIRGIN ISLANDS U.S.",
            "1345":"CAYMAN ISLANDS", "1441":"BERMUDA", "1473":"GRENADA", "1481":"GUERNSEY", "1534":"JERSEY", "1624":"ISLE OF MAN",
            "1649":"TURKS AND CAICOS ISLANDS", "1664":"MONTSERRAT", "1670":"NORTHERN MARIANA ISLANDS", "1671":"GUAM", "9999":"OTHERS",
            "6":"HOLY SEE (VATICAN CITY STATE)", "58":"VENEZUELA BOLIVARIAN REPUBLIC OF", "84":"VIET NAM", "260":"ZAMBIA", "263":"ZIMBABWE",
            "598":"URUGUAY", "678":"VANUATU", "681":"WALLIS AND FUTUNA", "967":"YEMEN", "998":"UZBEKISTAN", "1001":"ALAND ISLANDS",
            "1002":"BONAIRE SINT EUSTATIUS AND SABA", "1003":"BOUVET ISLAND", "1004":"FRENCH SOUTHERN TERRITORIES", "1005":"HEARD ISLAND AND MCDONALD ISLANDS",
            "1008":"SOUTH GEORGIA AND THE SOUTH SANDWICH ISLANDS", "1010":"ANTARCTICA", "1011":"PITCAIRN", "1013":"WESTERN SAHARA", "1016":"KOSOVO",
            "1721":"SINT MAARTEN (DUTCH PART)", "1758":"SAINT LUCIA", "1767":"DOMINICA", "1784":"SAINT VINCENT AND THE GRENADINES", "1787":"PUERTO RICO",
            "1809":"DOMINICAN REPUBLIC", "1868":"TRINIDAD AND TOBAGO", "1869":"SAINT KITTS AND NEVIS", "1876":"JAMAICA"
        }
        
        state_code = {
            "2":"Andhra Pradesh", "3":"Arunachal Pradesh", "4":"Assam", "5":"Bihar", "6":"Chandigarh", "8":"Daman and Diu", "9":"Delhi",
            "10":"Goa", "11":"Gujarat", "12":"Haryana", "13":"Himachal Pradesh", "14":"Jammu and Kashmir", "15":"Karnataka", "16":"Kerala",
            "17":"Lakshadweep", "18":"Madhya Pradesh", "19":"Maharashtra", "20":"Manipur", "22":"Mizoram", "23":"Nagaland", "24":"Odisha",
            "26":"Punjab", "27":"Rajasthan", "28":"Sikkim", "29":"Tamil Nadu", "30":"Tripura", "31":"Uttar Pradesh", "32":"West Bengal",
            "34":"Uttarakhand", "35":"Jharkhand", "36":"Telangana", "99":"Foreign", "37":"Ladakh", "1":"Andaman And Nicobar Islands",
            "7":"Dadra & Nagar Haveli", "21":"Meghalaya", "25":"Puducherry", "33":"Chhattisgarh"
        }

        #----------------------------------------BASIC USER IFNO----------------------------------------
        try:   
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
        
            if UserInfo_1.status_code == 200:
                createdTmstmp_encrypt = basic_info.get("createdTmstmp")
                if createdTmstmp_encrypt is not None:
                    createdTmstmpDate = datetime.fromtimestamp(createdTmstmp_encrypt / 1000)
                    createdTmstmp = createdTmstmpDate.strftime('%d-%m-%Y')
                    createdTmstmpTime = createdTmstmpDate.time()
                else:
                    createdTmstmp = "Null"
                    createdTmstmpTime = "Null"

                lastUpdatedTmstmp_encrypt = basic_info.get("lastUpdatedTmstmp")
                if lastUpdatedTmstmp_encrypt is not None:
                    lastUpdatedTmstmpDate = datetime.fromtimestamp(lastUpdatedTmstmp_encrypt / 1000)
                    lastUpdatedTmstmp = lastUpdatedTmstmpDate.strftime('%d-%m-%Y')
                    lastUpdatedTmstmpTime = lastUpdatedTmstmpDate.time()
                else:
                    lastUpdatedTmstmp = "Null"
                    lastUpdatedTmstmpTime = "Null"

                regStartDt_encrypt = basic_info.get("regStartDt")
                if regStartDt_encrypt is not None:
                    regStartDtDate = datetime.fromtimestamp(regStartDt_encrypt / 1000)
                    regStartDt = regStartDtDate.strftime('%d-%m-%Y')
                    regStartDtTime = regStartDtDate.time()
                else:
                    regStartDt = "Null"
                    regStartDtTime = "Null"  
                
                activationDt_encrypt = basic_info.get("activationDt")
                if activationDt_encrypt is not None:
                    activationDtDate = datetime.fromtimestamp(activationDt_encrypt / 1000)
                    activationDt = activationDtDate.strftime('%d-%m-%Y')
                    activationDtTime = activationDtDate.time()
                else:
                    activationDt = "Null"
                    activationDtTime = "Null"

                lastLoginTmstmp_encrypt = basic_info.get("lastLoginTmstmp")
                if lastLoginTmstmp_encrypt is not None:
                    lastLoginTmstmpDate = datetime.fromtimestamp(lastLoginTmstmp_encrypt / 1000)
                    lastLoginTmstmp = lastLoginTmstmpDate.strftime('%d-%m-%Y')
                    lastLoginTmstmpTime = lastLoginTmstmpDate.time()
                else:
                    lastLoginTmstmp = "Null"
                    lastLoginTmstmpTime = "Null"
                
                dscExpDt_encrypt = basic_info.get("dscExpDt")
                if dscExpDt_encrypt is not None:
                    dscExpDtDate = datetime.fromtimestamp(dscExpDt_encrypt / 1000)
                    dscExpDt = dscExpDtDate.strftime('%d-%m-%Y')
                    dscExpDtTime = dscExpDtDate.time()
                else:
                    dscExpDt = "Null"
                    dscExpDtTime = "Null"

                lastLogoutTmstmp_encrypt = basic_info.get("lastLogoutTmstmp")
                if lastLogoutTmstmp_encrypt is not None:
                    lastLogoutTmstmpDate = datetime.fromtimestamp(lastLogoutTmstmp_encrypt / 1000)
                    lastLogoutTmstmp = lastLogoutTmstmpDate.strftime('%d-%m-%Y')
                    lastLogoutTmstmpTime = lastLogoutTmstmpDate.time()
                else:
                    lastLogoutTmstmp = "Null"
                    lastLogoutTmstmpTime = "Null"

                priMobBelongsTo_encrypt = basic_info.get("priMobBelongsTo")
                if priMobBelongsTo_encrypt is not None:
                    priMobBelongsTo = No_Belongs_to.get(priMobBelongsTo_encrypt, "")
                else:
                    priMobBelongsTo = "Null"

                priEmailRelationId_encrypt = basic_info.get("priEmailRelationId")
                if priEmailRelationId_encrypt is not None:
                    priEmailRelationId = No_Belongs_to.get(priEmailRelationId_encrypt, "")
                else:
                    priEmailRelationId = "Null"

                secMobRelationId_encrypt = basic_info.get("secMobRelationId")
                if secMobRelationId_encrypt is not None:
                    secMobRelationId = No_Belongs_to.get(secMobRelationId_encrypt, "")
                else:
                    secMobRelationId = "Null"

                secEmailRelationId_encrypt = basic_info.get("secEmailRelationId")
                if secEmailRelationId_encrypt is not None:
                    secEmailRelationId = No_Belongs_to.get(secEmailRelationId_encrypt, "")
                else:
                    secEmailRelationId = "Null"

                contactResStatusCd_encrypt = basic_info.get("contactResStatusCd")
                if contactResStatusCd_encrypt is not None:
                    contactResStatusCd = resdiantial_status.get(contactResStatusCd_encrypt, "")
                else:
                    contactResStatusCd = "Null"

                residentialStatusCd_encrypt = basic_info.get("residentialStatusCd")
                if residentialStatusCd_encrypt is not None:
                    residentialStatusCd = resdiantial_status.get(residentialStatusCd_encrypt, "")
                else:
                    residentialStatusCd = "Null"

                stateCd_encrypt = basic_info.get("stateCd")
                if stateCd_encrypt is not None:
                    stateCd = state_code.get(stateCd_encrypt, "")
                else:
                    stateCd = "Null"

                countryCd_encrypt = basic_info.get("countryCd")
                if countryCd_encrypt is not None:
                    countryCd = country_code.get(countryCd_encrypt, "")
                else:
                    countryCd = "Null"

                aadhaarNum_encrypt = basic_info.get("aadhaarNum")
                if aadhaarNum_encrypt is not None:
                    aadhaarNum = base64.b64decode(aadhaarNum_encrypt).decode('utf-8')
                else:
                    aadhaarNum = "Null"

                user_basic_info_list = [{
                    "CA_reg_number":ca_registration_number,
                    "userId": basic_info.get("userId", "") or "Null",
                    "roleDesc": basic_info.get("roleDesc", "") or "Null",
                    "incorporateDate":basic_info.get("incorporateDate", "") or "Null",
                    "orgName":basic_info.get("orgName", "") or "Null",
                    "firstName": basic_info.get("firstName", "") or "Null",
                    "midName": basic_info.get("midName", "") or "Null",
                    "lastName": basic_info.get("lastName", "") or "Null",
                    "contactFirstName":basic_info.get("contactFirstName", "") or "Null",
                    "contactMiddleName":basic_info.get("contactMiddleName", "") or "Null",
                    "contactLastName":basic_info.get("contactLastName", "") or "Null",
                    "contactDesig":basic_info.get("contactDesig", "") or "Null",
                    "aadhaarNum": aadhaarNum,
                    "priMobileNum": basic_info.get("priMobileNum", "") or "Null",
                    "priMobBelongsTo":priMobBelongsTo,
                    "priEmailId": basic_info.get("priEmailId", "") or "Null",
                    "priEmailRelationId":priEmailRelationId,
                    "secMobileNum":basic_info.get("secMobileNum", "") or "Null",
                    "secMobRelationId":secMobRelationId,
                    "secEmailId":basic_info.get("secEmailId", "") or "Null",
                    "secEmailRelationId":secEmailRelationId,
                    "addrLine1Txt": basic_info.get("addrLine1Txt", "") or "Null",
                    "addrLine2Txt": basic_info.get("addrLine2Txt", "") or "Null",
                    "addrLine3Txt": basic_info.get("addrLine3Txt", "") or "Null",
                    "addrLine4Txt": basic_info.get("addrLine4Txt", "") or "Null",
                    "addrLine5Txt": basic_info.get("addrLine5Txt", "") or "Null",
                    "residentialStatusCd": residentialStatusCd,
                    "contactResStatusCd":contactResStatusCd,
                    "pinCd": basic_info.get("pinCd", "") or "Null",
                    "stateCd": stateCd,
                    "countryCd": countryCd,
                    "createdTmstmp": createdTmstmp,
                    "createdTmstmpTime":createdTmstmpTime,
                    "lastUpdatedTmstmp": lastUpdatedTmstmp,
                    "lastUpdatedTmstmpTime":lastUpdatedTmstmpTime,
                    "createdBy": basic_info.get("createdBy", "") or "Null",
                    "lastUpdatedBy": basic_info.get("lastUpdatedBy", "") or "Null",
                    "status": basic_info.get("status", "") or "Null",
                    "regStartDt":regStartDt,
                    "regStartDtTime":regStartDtTime,
                    "activationDt":activationDt,
                    "activationDtTime":activationDtTime,
                    "lastLoginTmstmp": lastLoginTmstmp,
                    "lastLoginTmstmpTime":lastLoginTmstmpTime,
                    "activationCode":basic_info.get("activationCode", "") or "Null",
                    "dscFlag":basic_info.get("dscFlag", "") or "Null",
                    "isMigrated":basic_info.get("isMigrated", "") or "Null",
                    "oldTranId":basic_info.get("oldTranId", "") or "Null",
                    "transactionNo": basic_info.get("transactionNo", "") or "Null",
                    "securedLogin":basic_info.get("securedLogin", "") or "Null",
                    "createdByUser": basic_info.get("createdByUser", "") or "Null",
                    "updatedByUser": basic_info.get("updatedByUser", "") or "Null",
                    "panStatus": basic_info.get("panStatus", "") or "Null",
                    "contactGender":basic_info.get("contactMiddleName", "") or "Null",
                    "userGender":basic_info.get("userGender", "") or "Null",
                    "dscExpDt":dscExpDt,
                    "dscExpDtTime":dscExpDtTime,
                    "dateOfBirth": basic_info.get("dob", "") or "Null",
                    "lastLogoutTmstmp":lastLogoutTmstmp,
                    "lastLogoutTmstmpTime":lastLogoutTmstmpTime,
                    "logoutCapturedFlg": basic_info.get("logoutCapturedFlg", "") or "Null",
                }]

                Information = pd.DataFrame(user_basic_info_list)
                Information.replace('\x00', '', regex=True, inplace=True)

                csv_file_path = basic_user_info
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'user_info'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE user_info (
                            CA_reg_number text,
                            userId text,
                            roleDesc text,
                            incorporateDate text,
                            orgName text,
                            firstName text,
                            midName text,
                            lastName text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            contactDesig text,
                            aadhaarNum text,
                            priMobileNum text,
                            priMobBelongsTo text,
                            priEmailId text,
                            priEmailRelationId text,
                            secMobileNum text,
                            secMobRelationId text,
                            secEmailId text,
                            secEmailRelationId text,
                            addrLine1Txt text,
                            addrLine2Txt text,
                            addrLine3Txt text,
                            addrLine4Txt text,
                            addrLine5Txt text,
                            residentialStatusCd text,
                            contactResStatusCd text,
                            pinCd text,
                            stateCd text,
                            countryCd text,
                            createdTmstmp text,
                            createdTmstmpTime text,
                            lastUpdatedTmstmp text,
                            lastUpdatedTmstmpTime text,
                            createdBy text,
                            lastUpdatedBy text,
                            status text,
                            regStartDt text,
                            regStartDtTime text,
                            activationDt text,
                            activationDtTime text,
                            lastLoginTmstmp text,
                            lastLoginTmstmpTime text,
                            activationCode text,
                            dscFlag text,
                            isMigrated text,
                            oldTranId text,
                            transactionNo text,
                            securedLogin text,
                            createdByUser text,
                            updatedByUser text,
                            panStatus text,
                            contactGender text,
                            userGender text,
                            dscExpDt text,
                            dscExpDtTime text,
                            dateOfBirth text,
                            lastLogoutTmstmp text,
                            lastLogoutTmstmpTime text,
                            logoutCapturedFlg text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'user_info' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"USER INFORMATION: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"USER INFORMATION: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO user_info (
                                    CA_reg_number, userId, roleDesc, incorporateDate, orgName, 
                                    firstName, midName, lastName, contactFirstName, contactMiddleName, contactLastName, 
                                    contactDesig, aadhaarNum, priMobileNum, priMobBelongsTo, priEmailId, 
                                    priEmailRelationId, secMobileNum, secMobRelationId, secEmailId, 
                                    secEmailRelationId, addrLine1Txt, addrLine2Txt, addrLine3Txt, addrLine4Txt, 
                                    addrLine5Txt, residentialStatusCd, contactResStatusCd, pinCd, stateCd, 
                                    countryCd, createdTmstmp, createdTmstmpTime, lastUpdatedTmstmp, lastUpdatedTmstmpTime, 
                                    createdBy, lastUpdatedBy, status, regStartDt, regStartDtTime, activationDt, 
                                    activationDtTime, lastLoginTmstmp, lastLoginTmstmpTime, activationCode, 
                                    dscFlag, isMigrated, oldTranId, transactionNo, securedLogin, createdByUser, 
                                    updatedByUser, panStatus, contactGender, userGender, dscExpDt, dscExpDtTime, 
                                    dateOfBirth, lastLogoutTmstmp, lastLogoutTmstmpTime, logoutCapturedFlg
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['CA_reg_number'], row['userId'], row['roleDesc'], row['incorporateDate'], row['orgName'],
                                    row['firstName'], row['midName'], row['lastName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'],
                                    row['contactDesig'], row['aadhaarNum'], row['priMobileNum'], row['priMobBelongsTo'], row['priEmailId'],
                                    row['priEmailRelationId'], row['secMobileNum'], row['secMobRelationId'], row['secEmailId'],
                                    row['secEmailRelationId'], row['addrLine1Txt'], row['addrLine2Txt'], row['addrLine3Txt'], row['addrLine4Txt'],
                                    row['addrLine5Txt'], row['residentialStatusCd'], row['contactResStatusCd'], row['pinCd'], row['stateCd'],
                                    row['countryCd'], row['createdTmstmp'], row['createdTmstmpTime'], row['lastUpdatedTmstmp'], row['lastUpdatedTmstmpTime'],
                                    row['createdBy'], row['lastUpdatedBy'], row['status'], row['regStartDt'], row['regStartDtTime'], row['activationDt'],
                                    row['activationDtTime'], row['lastLoginTmstmp'], row['lastLoginTmstmpTime'], row['activationCode'],
                                    row['dscFlag'], row['isMigrated'], row['oldTranId'], row['transactionNo'], row['securedLogin'], row['createdByUser'],
                                    row['updatedByUser'], row['panStatus'], row['contactGender'], row['userGender'], row['dscExpDt'], row['dscExpDtTime'],
                                    row['dateOfBirth'], row['lastLogoutTmstmp'], row['lastLogoutTmstmpTime'], row['logoutCapturedFlg']
                                ))

                                conn_ef_db.commit()
                                print(f"User Info: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("USER INFORMATION: CSV CREATED SUCCESSFULLY")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO user_info (
                            CA_reg_number, userId, roleDesc, incorporateDate, orgName, 
                            firstName, midName, lastName, contactFirstName, contactMiddleName, contactLastName, 
                            contactDesig, aadhaarNum, priMobileNum, priMobBelongsTo, priEmailId, 
                            priEmailRelationId, secMobileNum, secMobRelationId, secEmailId, 
                            secEmailRelationId, addrLine1Txt, addrLine2Txt, addrLine3Txt, addrLine4Txt, 
                            addrLine5Txt, residentialStatusCd, contactResStatusCd, pinCd, stateCd, 
                            countryCd, createdTmstmp, createdTmstmpTime, lastUpdatedTmstmp, lastUpdatedTmstmpTime, 
                            createdBy, lastUpdatedBy, status, regStartDt, regStartDtTime, activationDt, 
                            activationDtTime, lastLoginTmstmp, lastLoginTmstmpTime, activationCode, 
                            dscFlag, isMigrated, oldTranId, transactionNo, securedLogin, createdByUser, 
                            updatedByUser, panStatus, contactGender, userGender, dscExpDt, dscExpDtTime, 
                            dateOfBirth, lastLogoutTmstmp, lastLogoutTmstmpTime, logoutCapturedFlg
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                        """, (
                            row['CA_reg_number'], row['userId'], row['roleDesc'], row['incorporateDate'], row['orgName'],
                            row['firstName'], row['midName'], row['lastName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'],
                            row['contactDesig'], row['aadhaarNum'], row['priMobileNum'], row['priMobBelongsTo'], row['priEmailId'],
                            row['priEmailRelationId'], row['secMobileNum'], row['secMobRelationId'], row['secEmailId'],
                            row['secEmailRelationId'], row['addrLine1Txt'], row['addrLine2Txt'], row['addrLine3Txt'], row['addrLine4Txt'],
                            row['addrLine5Txt'], row['residentialStatusCd'], row['contactResStatusCd'], row['pinCd'], row['stateCd'],
                            row['countryCd'], row['createdTmstmp'], row['createdTmstmpTime'], row['lastUpdatedTmstmp'], row['lastUpdatedTmstmpTime'],
                            row['createdBy'], row['lastUpdatedBy'], row['status'], row['regStartDt'], row['regStartDtTime'], row['activationDt'],
                            row['activationDtTime'], row['lastLoginTmstmp'], row['lastLoginTmstmpTime'], row['activationCode'],
                            row['dscFlag'], row['isMigrated'], row['oldTranId'], row['transactionNo'], row['securedLogin'], row['createdByUser'],
                            row['updatedByUser'], row['panStatus'], row['contactGender'], row['userGender'], row['dscExpDt'], row['dscExpDtTime'],
                            row['dateOfBirth'], row['lastLogoutTmstmp'], row['lastLogoutTmstmpTime'], row['logoutCapturedFlg']
                        ))

                        conn_ef_db.commit()
                        print(f"User Info: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()

        except Exception as e:
            print(f"{username} An error occurred for basic information: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_user_info
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"BASIC INFORMATION: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"BASIC INFORMATION: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"BASIC INFORMATION: ERROR USERS CSV CREATED.")
        #----------------------------------------BASIC INFO END----------------------------------------

        #----------------------------------------REFUND/DEMAND START----------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)

            payload_refund_demand = json.dumps({
            "header": {
                        "formName": "FO-013-DSBRD"
                    },
                    "entityNum": response_entity,
                    "assmentYear": str(last_year)
            })

            headers_refund_demand = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'NA'
            }

            payload_refund_demand_2 = json.dumps({
                "header": {
                        "formName": "FO-002-LOGIN"
                    },
                    "serviceName": "itrStagesSevice",
                    "pan": response_entity,
                    "ay": str(last_year)
            })

            headers_refund_demand_2 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'NA'
            }

            response_refund_demand = requests.request("POST", refund_demand_url, headers = headers_refund_demand, data = payload_refund_demand)
            response_refund_demand_info = requests.request("POST", refund_demand_info, headers = headers_refund_demand_2, data = payload_refund_demand_2)
            time.sleep(3)

            refund_demand = json.loads(response_refund_demand.text)
            refund_demand_information = json.loads(response_refund_demand_info.text)
            time.sleep(3)

            if response_refund_demand.status_code == 200 and response_refund_demand_info.status_code == 200:
                try:
                    status_value = refund_demand['status']
                    refund_amount=refund_demand['processedRefundAmt']
                except KeyError:
                    status_value = "No Data Present"
                    refund_amount="0"

                refund_demand_header = [{
                "CA_reg_number": ca_registration_number,
                "userId": basic_info.get("userId", ""),
                "roleDesc": basic_info.get("roleDesc", ""),
                "orgName":basic_info.get("orgName", "")or "Null",
                "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                "contactLastName":basic_info.get("contactLastName", "")or "Null",
                "firstName": basic_info.get("firstName", ""),
                "midName": basic_info.get("midName", ""),
                "lastName": basic_info.get("lastName", ""),
                "ReturnFiledOn":refund_demand_information.get("ackDt", ""),
                "ReturnVarifiedOn":refund_demand_information.get("ackDt", ""),
                "VarificationStatus":refund_demand_information.get("verStatus", ""),
                "ReturnProcessing":refund_demand_information.get("processingDate", ""),
                "ProcessingComplition":refund_demand_information.get("completedDate", ""),
                "Year": last_year,
                "refund_amount": refund_amount,
                "status": status_value
                }]

                Information = pd.DataFrame(refund_demand_header)
                Information.replace('\x00', '', regex=True, inplace=True)

                csv_file_path = refund_demand_path
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'refund_demand'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE refund_demand (
                            CA_reg_number text,
                            userId text,
                            roleDesc text,
                            orgName text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            firstName text,
                            midName text,
                            lastName text,
                            ReturnFiledOn text,
                            ReturnVarifiedOn text,
                            VarificationStatus text,
                            ReturnProcessing text,
                            ProcessingComplition text,
                            Year text,
                            refund_amount text,
                            status text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'refund_demand' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"REFUND/DEMAND: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"REFUND/DEMAND: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO refund_demand (
                                    CA_reg_number, userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                    firstName, midName, lastName, ReturnFiledOn, ReturnVarifiedOn, 
                                    VarificationStatus, ReturnProcessing, ProcessingComplition, 
                                    Year, refund_amount, status
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['CA_reg_number'], row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                    row['firstName'],  row['midName'], row['lastName'], row['ReturnFiledOn'], row['ReturnVarifiedOn'], 
                                    row['VarificationStatus'], row['ReturnProcessing'], row['ProcessingComplition'], 
                                    row['Year'], row['refund_amount'], row['status']
                                ))

                                conn_ef_db.commit()
                                print(f"Refund/Demand: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()
                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("REFUND/DEMAND CSV CREATED SUCCESSFULLY.")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO refund_demand (
                            CA_reg_number, userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                            firstName, midName, lastName, ReturnFiledOn, ReturnVarifiedOn, 
                            VarificationStatus, ReturnProcessing, ProcessingComplition, 
                            Year, refund_amount, status
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                        """, (
                            row['CA_reg_number'], row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                            row['firstName'],  row['midName'], row['lastName'], row['ReturnFiledOn'], row['ReturnVarifiedOn'], 
                            row['VarificationStatus'], row['ReturnProcessing'], row['ProcessingComplition'], 
                            row['Year'], row['refund_amount'], row['status']
                        ))

                        conn_ef_db.commit()
                        print(f"Refund/Demand: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()
        
        except Exception as e:
            print(f"{username} An error occurred for refund/demand: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_refund_demand
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"REFUND/DEMAND: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"REFUND/DEMAND: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"REFUND/DEMAND: ERROR USERS CSV CREATED.")
        #----------------------------------------REFUND/DEMAND END------------------------------------------

        #----------------------------------------TAX DEPOSIT START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)

            payload_taxdeposit = json.dumps({
            "serviceName": "taxDepositService",
            "pan": response_entity,
            "ay": 2024
            })

            headers_taxdeposit = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            Tax_Deposit = requests.request("POST", tax_deposit_url, headers = headers_taxdeposit, data = payload_taxdeposit)
            time.sleep(3)

            basic_info = json.loads(UserInfo_1.text)
            tax_deposit = json.loads(Tax_Deposit.text)
            time.sleep(3)

            if UserInfo_1.status_code == 200 and Tax_Deposit.status_code == 200:
                data_list = []
                for item in tax_deposit:
                    satAmt = item.get("satAmt", "")
                    atAmt = item.get("atAmt", "")
                    tdsAmt = item.get("tdsAmt", "")
                    tcsAmt = item.get("tcsAmt", "")
                    ay = item.get("ay", "")
                    
                    user_basic_info_list = {
                        "userId": basic_info.get("userId", "") or "Null",
                        "orgName":basic_info.get("orgName", "") or "Null",
                        "contactFirstName":basic_info.get("contactFirstName", "") or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "") or "Null",
                        "contactLastName":basic_info.get("contactLastName", "") or "Null",
                        "firstName": basic_info.get("firstName", "") or "Null",
                        "midName": basic_info.get("midName", "") or "Null",
                        "lastName": basic_info.get("lastName", "") or "Null",
                        "satAmt" : satAmt,
                        "atAmt" : atAmt,
                        "tdsAmt" : tdsAmt,
                        "tcsAmt" : tcsAmt,
                        "ay" : ay
                    }
                    data_list.append(user_basic_info_list)

                Information = pd.DataFrame(data_list)
                Information.replace('\x00', '', regex=True, inplace=True)

                csv_file_path = tax_deposit_path
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'tax_deposit'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE tax_deposit (
                            userId text,
                            orgName text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            firstName text,
                            midName text,
                            lastName text,
                            satAmt text,
                            atAmt text,
                            tdsAmt text,
                            tcsAmt text,
                            ay text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'tax_deposit' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"TAX DEPOSIT: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"TAX DEPOSIT: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO tax_deposit (
                                    userId, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                    firstName, midName, lastName, satAmt, atAmt, tdsAmt, tcsAmt, ay
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s;
                                """, (
                                    row['userId'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                    row['firstName'],  row['midName'], row['lastName'], row['satAmt'], row['atAmt'], row['tdsAmt'], row['tcsAmt'], 
                                    row['ay']
                                ))

                                conn_ef_db.commit()
                                print(f"Tax Deposit: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("USER INFO saved to CSV successfully.")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO tax_deposit (
                            userId, orgName, contactFirstName, contactMiddleName, contactLastName, 
                            firstName, midName, lastName, satAmt, atAmt, tdsAmt, tcsAmt, ay
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s;
                        """, (
                            row['userId'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                            row['firstName'],  row['midName'], row['lastName'], row['satAmt'], row['atAmt'], row['tdsAmt'], row['tcsAmt'], 
                            row['ay']
                        ))

                        conn_ef_db.commit()
                        print(f"Tax Deposit: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()

        except Exception as e:
            print(f"{username} An error occurred for tax deposit: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_tax_deposit
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"TAX DEPOSIT: ERROR USER RECORD{row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"TAX DEPOIST: ERROR USER RECORD{row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"TAX DEPOSIT: ERROR USERS CSV CREATED.")
        #----------------------------------------TAX DEPOSIT END------------------------------------------

        #----------------------------------------RECENT FILED RETURN START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_recentfilereturn = json.dumps({
            "pan": response_entity
            })

            headers_recentfilereturn = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            Recent_File_Return = requests.request("POST", recent_filed_return_url, headers = headers_recentfilereturn, data = payload_recentfilereturn)
            time.sleep(3)

            basic_info = json.loads(UserInfo_1.text)
            recent_file_return = json.loads(Recent_File_Return.text)
            time.sleep(3)

            if UserInfo_1.status_code == 200 and Recent_File_Return.status_code == 200:
                if recent_file_return.get("lastThreeYearsReturn"):
                    recentfile_return = recent_file_return.get("lastThreeYearsReturn")

                    data_list = []
                    for item in recentfile_return:
                        AssesmentYear = item.get("AssesmentYear", "")
                        TaxableIncome = item.get("TaxableIncome", "")
                        TaxLiability = item.get("TaxLiability", "")
                        TaxDeposited = item.get("TaxDeposited", "")
                        
                        user_basic_info_list = {
                        "userId": basic_info.get("userId", ""),
                        "roleDesc": basic_info.get("roleDesc", ""),
                        "orgName":basic_info.get("orgName", "")or "Null",
                        "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                        "contactLastName":basic_info.get("contactLastName", "")or "Null",
                        "firstName": basic_info.get("firstName", ""),
                        "midName": basic_info.get("midName", ""),
                        "lastName": basic_info.get("lastName", ""),
                        "AssesmentYear" : AssesmentYear,
                        "TaxableIncome" : TaxableIncome,
                        "TaxLiability" : TaxLiability,
                        "TaxDeposited" : TaxDeposited
                        }
                        data_list.append(user_basic_info_list)

                    Information = pd.DataFrame(data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                
                    csv_file_path = recent_filed_return
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'recentfiledreturn'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE recentfiledreturn (
                                userId text,
                                roleDesc text,
                                orgName text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                AssesmentYear text,
                                TaxableIncome text,
                                TaxLiability text,
                                TaxDeposited text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'recentfiledreturn' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['userId'].tolist()

                            for index, row in Information.iterrows():
                                if row['userId'] in existing_users:
                                    print(f"RECENT FILED RETURN: {row['userId']} ALREADY EXISTED.")

                                else:
                                        with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                            Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                            print(f"RECENT FILED RETURN: {row['userId']} INSERTED SUCCESSFULLY.")

                                        curr_ef_db.execute("""
                                        INSERT INTO recentfiledreturn (
                                            userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                            firstName, midName, lastName, AssesmentYear, TaxableIncome, TaxLiability, TaxDeposited
                                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                        """, (
                                            row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                            row['firstName'],  row['midName'], row['lastName'], row['AssesmentYear'], row['TaxableIncome'], 
                                            row['TaxLiability'], row['TaxDeposited']
                                        ))

                                        conn_ef_db.commit()
                                        print(f"Recent Filed Return: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                        print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("RECENT FILED RETURN: CSV CREATED SUCCESSFULLY.")
                        
                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO recentfiledreturn (
                                userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, AssesmentYear, TaxableIncome, TaxLiability, TaxDeposited
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                row['firstName'],  row['midName'], row['lastName'], row['AssesmentYear'], row['TaxableIncome'], 
                                row['TaxLiability'], row['TaxDeposited']
                            ))

                            conn_ef_db.commit()
                            print(f"Recent Filed Return: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()

        except Exception as e:
            print(f"{username} An error occurred for recent filed return: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_recent_file_return
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"RECENT FILED RETURN: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"RECENT FILED RETURN: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"RECENT FILED RETURN: ERROR USERS CSV CREATED.")
        #----------------------------------------RECENT FILED RETURN END------------------------------------------

        #----------------------------------------RECENT FORM FILED START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_recentformsfiled = json.dumps({
            "serviceName": "viewFiledForms",
            "entityNum": response_entity
            })

            headers_recentformsfiled = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            Recent_Forms_Filed = requests.request("POST", recent_forms_filed_url, headers = headers_recentformsfiled, data = payload_recentformsfiled)
            time.sleep(3)

            basic_info = json.loads(UserInfo_1.text)
            recent_forms_filed = json.loads(Recent_Forms_Filed.text)
            time.sleep(3)

            if UserInfo_1.status_code == 200 and Recent_Forms_Filed.status_code == 200:
                forms = recent_forms_filed.get("forms")
                if forms:
                    ref_year = recent_forms_filed["forms"][0]["refYear"]
                    ack_date = recent_forms_filed["forms"][0]["ackDate"]
                    
                    data_list = []
                    for item in forms:
                        formName = item.get("formName", "")
                        formShortName = item.get("formShortName", "")
                        formDesc = item.get("formDesc", "")
                        formCd = item.get("formCd", "")
                        fillingCount = item.get("fillingCount", "")
                        refYearType = item.get("refYearType", "")
                        mode = item.get("mode", "")
                        formNameHindi = item.get("formNameHindi", "")
                        formShortNameHindi = item.get("formShortNameHindi", "")

                        
                        user_basic_info_list = {
                        "userId": basic_info.get("userId", ""),
                        "roleDesc": basic_info.get("roleDesc", ""),
                        "orgName":basic_info.get("orgName", "")or "Null",
                        "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                        "contactLastName":basic_info.get("contactLastName", "")or "Null",
                        "firstName": basic_info.get("firstName", ""),
                        "midName": basic_info.get("midName", ""),
                        "lastName": basic_info.get("lastName", ""),
                        "userType" : recent_forms_filed.get("userType", "")or "Null",
                        "submitUserId" : recent_forms_filed.get("submitUserId", "") or "Null",
                        "formCount" : recent_forms_filed.get("formCount", ""),
                        "eriPan" : recent_forms_filed.get("eriPan", "") or "Null",
                        "mode" : mode,
                        "formName" : formName,
                        "formShortName" : formShortName,
                        "formDesc" : formDesc,
                        "formCd" : formCd,
                        "fillingCount" : fillingCount,
                        "refYearType" : refYearType,
                        "ref_year" : ref_year,
                        "ack_date" : ack_date,
                        "formNameHindi" : formNameHindi,
                        "formShortNameHindi":formShortNameHindi
                        }
                        data_list.append(user_basic_info_list)

                    Information = pd.DataFrame(data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                
                    csv_file_path = recent_form_filed_path
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'recent_form_filed'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE recent_form_filed (
                                userId text,
                                roleDesc text,
                                orgName text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                userType text,
                                submitUserId text,
                                formCount text,
                                eriPan text,
                                mode text,
                                formName text,
                                formShortName text,
                                formDesc text,
                                formCd text,
                                fillingCount text,
                                refYearType text,
                                ref_year text,
                                ack_date text,
                                formNameHindi text,
                                formShortNameHindi text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'recent_form_filed' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['userId'].tolist()

                            for index, row in Information.iterrows():
                                if row['userId'] in existing_users:
                                    print(f"RECENT FORM FILED: {row['userId']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"RECENT FORM FILED: {row['userId']} INSERTED SUCCESSFULLY.")
                                    
                                    curr_ef_db.execute("""
                                    INSERT INTO recent_form_filed (
                                        userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, userType, submitUserId, formCount, 
                                        eriPan, mode, formName, formShortName, formDesc, formCd, 
                                        fillingCount, refYearType, ref_year, ack_date, formNameHindi, formShortNameHindi
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['userId'], row['roleDesc'],  row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                        row['firstName'],  row['midName'], row['lastName'], row['userType'], row['submitUserId'], row['formCount'],
                                        row['eriPan'], row['mode'], row['formName'], row['formShortName'], row['formDesc'], row['formCd'],
                                        row['fillingCount'], row['refYearType'], row['ref_year'], row['ack_date'], row['formNameHindi'], row['formShortNameHindi']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Recent Form Filed: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("RECENT FORM FILED: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO recent_form_filed (
                                userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, userType, submitUserId, formCount, 
                                eriPan, mode, formName, formShortName, formDesc, formCd, 
                                fillingCount, refYearType, ref_year, ack_date, formNameHindi, formShortNameHindi
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s %s, %s, %s, %s, %s, %s);
                            """, (
                                row['userId'], row['roleDesc'],  row['orgName'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                row['firstName'],  row['midName'], row['lastName'], row['userType'], row['submitUserId'], row['formCount'],
                                row['eriPan'], row['mode'], row['formName'], row['formShortName'], row['formDesc'], row['formCd'],
                                row['fillingCount'], row['refYearType'], row['ref_year'], row['ack_date'], row['formNameHindi'], row['formShortNameHindi']
                            ))

                            conn_ef_db.commit()
                            print(f"Recent Form Filed: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()
        
        except Exception as e:
            print(f"{username} An error occurred for recent form filed: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_recent_form_filed
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"RECENT FORM FILED: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"RECENT FORM FILED: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"RECENT FORM FILED: ERROR USERS CSV CREATED.")
        #----------------------------------------RECENT FORM FILED END------------------------------------------

        #----------------------------------------ACTIVE BANK START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)

            payload_bankInfo = json.dumps({
            "entityNum": response_entity,
            "serviceName": "myBankAccountService",
            "header": {
                "formName": "FO-054-PBACC"
            }
            })

            headers_bankInfo = {
            'Accept': 'application/json, text/plain, */*',
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'myBankAccountService',
            }

            active_bank_response = requests.request("POST", bank_info_url, headers = headers_bankInfo, data = payload_bankInfo)
            time.sleep(3)

            active_bank_info = json.loads(active_bank_response.text)

            if active_bank_response.status_code == 200:
                time.sleep(3)
                if active_bank_info.get('activeBank'):
                    active_bank = active_bank_info.get('activeBank', [])
                
                    data_list = []
                    for entry in active_bank:
                        encrypt_acc_no = entry.get("bankAcctNum")
                        acc_no = base64.b64decode(encrypt_acc_no).decode('utf-8')
                            
                        bank_info_list = {
                        'entity':entry.get("entityNum", ""),
                        'name': entry.get("nameAsPerBank", ""),
                        "orgName":basic_info.get("orgName", "")or "Null",
                        "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                        "contactLastName":basic_info.get("contactLastName", "")or "Null",
                        "firstName": basic_info.get("firstName", "") or "Null",
                        "midName": basic_info.get("midName", "") or "Null",
                        "lastName": basic_info.get("lastName", "") or "Null",
                        'bank_account_no':acc_no,
                        'ifsc_code':entry.get("ifscCd", "") or "Null",
                        'bank_name':entry.get("bankName", "") or "Null",
                        'branch_texts':entry.get("bankBrnchTxt", "") or "Null",
                        'role':entry.get("role", "") or "Null",
                        }
                        data_list.append(bank_info_list)

                    Information = pd.DataFrame(data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                            
                    csv_file_path = active_bank_path
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'active_bank'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE active_bank (
                                entity text,
                                name text,
                                orgName text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                bank_account_no text,
                                ifsc_code text,
                                bank_name text,
                                branch_texts text,
                                role text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'active_bank' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['entity'].tolist()

                            for index, row in Information.iterrows():
                                if row['entity'] in existing_users:
                                    print(f"ACTIVE BANK: {row['entity']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"ACTIVE BANK: {row['entity']} INSERTED SUCCESSFULLY.")

                                    curr_ef_db.execute("""
                                    INSERT INTO active_bank (
                                        entity, name, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, bank_account_no, ifsc_code, bank_name, branch_texts, role
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['entity'], row['name'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], 
                                        row['contactLastName'], row['firstName'], row['midName'], row['lastName'], row['bank_account_no'], 
                                        row['ifsc_code'], row['bank_name'], row['branch_texts'], row['role']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Active Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("ACTIVE BANK INFORMATIN: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO active_bank (
                                entity, name, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, bank_account_no, ifsc_code, bank_name, branch_texts, role
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['entity'], row['name'], row['orgName'], row['contactFirstName'], row['contactMiddleName'], 
                                row['contactLastName'], row['firstName'], row['midName'], row['lastName'], row['bank_account_no'], 
                                row['ifsc_code'], row['bank_name'], row['branch_texts'], row['role']
                            ))

                            conn_ef_db.commit()
                            print(f"Active Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()
        
        except Exception as e:
            print(f"{username} An error occurred for active bank: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_active_bank
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"ACTIVE BANK: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"ACTIVE BANK: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"ACTIVE BANK: ERROR USERS CSV CREATED.")
        #----------------------------------------ACTIVE BANK END------------------------------------------

        #----------------------------------------IN-ACTIVR BANK START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)

            payload_inactiveBank = json.dumps({
            "entityNum": response_entity,
            "serviceName": "myBankAccountService",
            "header": {
                "formName": "FO-054-PBACC"
                }
            })

            headers_inactiveBank = {
            'Accept': 'application/json, text/plain, */*',
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'myBankAccountService',
            }

            inactive_bank_response = requests.request("POST", inactive_bank_url, headers = headers_inactiveBank, data = payload_inactiveBank)
            time.sleep(3)

            inactive_bank_info = json.loads(inactive_bank_response.text)
            time.sleep(3)

            if inactive_bank_response.status_code == 200:
                if inactive_bank_info.get('inActiveBank'):
                    inactive_bank = inactive_bank_info.get("inActiveBank", [])

                    data_list = []
                    for entry in inactive_bank:
                        encrypt_acc_no = entry.get("bankAcctNum")
                        acc_no = base64.b64decode(encrypt_acc_no).decode('utf-8')

                        bank_info_list = {
                        'entity':entry.get("entityNum", ""),
                        "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                        "contactLastName":basic_info.get("contactLastName", "")or "Null",
                        "firstName": basic_info.get("firstName", ""),
                        "midName": basic_info.get("midName", ""),
                        "lastName": basic_info.get("lastName", ""),
                        'bank_account_no':acc_no,
                        'ifsc_code':entry.get("ifscCd", ""),
                        'bank_name':entry.get("bankName", ""),
                        'branch_texts':entry.get("bankBrnchTxt", ""),
                        'error':entry.get("errorCd", ""),
                        'role':entry.get("role", ""),
                        'userAction':entry.get("userAction", "")
                        }                    
                        data_list.append(bank_info_list)
                
                    Information = pd.DataFrame(data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                            
                    csv_file_path = inactive_bank_path
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'inactive_bank'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE inactive_bank (
                                entity text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                bank_account_no text,
                                ifsc_code text,
                                bank_name text,
                                branch_texts text,
                                error text,
                                role text,
                                userAction text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'inactive_bank' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['entity'].tolist()

                            for index, row in Information.iterrows():
                                if row['entity'] in existing_users:
                                    print(f"INACTIVE BANK: {row['entity']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"INACTIVE BANK: {row['entity']} INSERTED SUCCESSFULLY.")

                                    curr_ef_db.execute("""
                                    INSERT INTO inactive_bank (
                                        entity, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, bank_account_no, ifsc_code, 
                                        bank_name, branch_texts, error, role, userAction
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['entity'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                        row['firstName'],  row['midName'], row['lastName'], row['bank_account_no'], row['ifsc_code'],
                                        row['bank_name'], row['branch_texts'], row['error'], row['role'], row['userAction']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"InActive Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("INACTIVE BANK INFORMATION: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO inactive_bank (
                                entity, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, bank_account_no, ifsc_code, 
                                bank_name, branch_texts, error, role, userAction
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['entity'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                row['firstName'],  row['midName'], row['lastName'], row['bank_account_no'], row['ifsc_code'],
                                row['bank_name'], row['branch_texts'], row['error'], row['role'], row['userAction']
                            ))

                            conn_ef_db.commit()
                            print(f"InActive Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()
        
        except Exception as e:
            print(f"{username} An error occurred for inactive bank: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_inactive_bank
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"INACTIVE BANK: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"INACTIVE BANK: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"INACTIVE BANK: ERROR USERS CSV CREATED.")
        #----------------------------------------IN-ACTIVR BANK END------------------------------------------

        #----------------------------------------FAILED BANK START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)

            payload_failedBank = json.dumps({
            "entityNum": response_entity,
            "serviceName": "myBankAccountService",
            "header": {
                "formName": "FO-054-PBACC"
                }
            })

            headers_failedBank = {
            'Accept': 'application/json, text/plain, */*',
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'myBankAccountService',
            }

            failed_bank_response = requests.request("POST", failed_bank_url, headers = headers_failedBank, data = payload_failedBank)
            time.sleep(3)

            failed_bank_info = json.loads(failed_bank_response.text)

            if failed_bank_response.status_code == 200:
                time.sleep(3)
                if failed_bank_info.get('failedBank'):
                    failed_bank = failed_bank_info.get("failedBank", [])

                    data_list = []
                    for entry in failed_bank:
                        encrypt_acc_no = entry.get("bankAcctNum")
                        acc_no = base64.b64decode(encrypt_acc_no).decode('utf-8')

                        bank_info_list = {
                        'entity':entry.get("entityNum", ""),
                        "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                        "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                        "contactLastName":basic_info.get("contactLastName", "")or "Null",
                        "firstName": basic_info.get("firstName", ""),
                        "midName": basic_info.get("midName", ""),
                        "lastName": basic_info.get("lastName", ""),
                        'bank_account_no':acc_no,
                        'ifsc_code':entry.get("ifscCd", ""),
                        'bank_name':entry.get("bankName", ""),
                        'branch_texts':entry.get("bankBrnchTxt", ""),
                        'error':entry.get("errorCd", ""),
                        'role':entry.get("role", ""),
                        'userAction':entry.get("userAction", "")
                        }
                        data_list.append(bank_info_list)
                
                    Information = pd.DataFrame(data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                            
                    csv_file_path = failed_bank_path
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'failed_bank'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE failed_bank (
                                entity text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                bank_account_no text,
                                ifsc_code text,
                                bank_name text,
                                branch_texts text,
                                error text,
                                role text,
                                userAction text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'failed_bank' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['entity'].tolist()

                            for index, row in Information.iterrows():
                                if row['entity'] in existing_users:
                                    print(f"FAILED BANK: {row['entity']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FAILED BANK: {row['entity']} INSERTED SUCCESSFULLY.")

                                    curr_ef_db.execute("""
                                    INSERT INTO failed_bank (
                                        entity, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, bank_account_no, ifsc_code, 
                                        bank_name, branch_texts, error, role, userAction
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['entity'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                        row['firstName'],  row['midName'], row['lastName'], row['bank_account_no'], row['ifsc_code'],
                                        row['bank_name'], row['branch_texts'], row['error'], row['role'], row['userAction']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Failed Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("FAILED BANK INFORMATION: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO failed_bank (
                                entity, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, bank_account_no, ifsc_code, 
                                bank_name, branch_texts, error, role, userAction
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['entity'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                row['firstName'],  row['midName'], row['lastName'], row['bank_account_no'], row['ifsc_code'],
                                row['bank_name'], row['branch_texts'], row['error'], row['role'], row['userAction']
                            ))

                            conn_ef_db.commit()
                            print(f"Failed Bank: {row['entity']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()
        
        except Exception as e:
            print(f"{username} An error occurred for failed bank: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_failed_bank
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"FAILED BANK: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"FAILED BANK: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"FAILED BANK: ERROR USERS CSV CREATED.")
        #----------------------------------------FAILED BANK END------------------------------------------

        #----------------------------------------JURISDCITION DETAILS START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }
            
            payload_juris = json.dumps({
            "serviceName": "jurisdictionDetailsService"
            })

            headers_juris = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'jurisdictionDetailsService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            juris = requests.request("POST", jurisdiction_url, headers = headers_juris, data = payload_juris)
            time.sleep(3)

            basic_info = json.loads(UserInfo_1.text)
            juris_info = json.loads(juris.text)
            time.sleep(3)

            if UserInfo_1.status_code == 200 and juris.status_code == 200:
                user_basic_info_list = [{
                "CA_reg_number":ca_registration_number,
                "ENTITY":response_entity,
                "userId": basic_info.get("userId", ""),
                "roleDesc": basic_info.get("roleDesc", ""),
                "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                "contactLastName":basic_info.get("contactLastName", "")or "Null",
                "firstName": basic_info.get("firstName", ""),
                "midName": basic_info.get("midName", ""),
                "lastName": basic_info.get("lastName", ""),
                "areaCd":juris_info.get("areaCd",""),
                "areaDesc":juris_info.get("areaDesc",""),
                "aoType":juris_info.get("aoType",""),
                "rangeCd":juris_info.get("rangeCd",""),
                "aoNo":juris_info.get("aoNo",""),   
                "aoPplrName":juris_info.get("aoPplrName",""),
                "aoEmailId":juris_info.get("aoEmailId",""),
                "aoBldgId":juris_info.get("aoBldgId",""),
                "aoBldgDesc":juris_info.get("aoBldgDesc",""),
                }]

                Information = pd.DataFrame(user_basic_info_list)
                Information.replace('\x00', '', regex=True, inplace=True)
                    
                csv_file_path = juris_diction
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'juris_diction_info'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE juris_diction_info (
                            CA_reg_number text,
                            ENTITY text,
                            userId text,
                            roleDesc text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            firstName text,
                            midName text,
                            lastName text,
                            areaCd text,
                            areaDesc text,
                            aoType text,
                            rangeCd text,
                            aoNo text,
                            aoPplrName text,
                            aoEmailId text,
                            aoBldgId text,
                            aoBldgDesc text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'juris_diction_info' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"JURISDCITION DETAILS: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"JURISDCITION DETAILS: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO juris_diction_info (
                                    CA_reg_number, ENTITY, userId, roleDesc, contactFirstName, contactMiddleName, contactLastName, 
                                    firstName, midName, lastName, areaCd, areaDesc, 
                                    aoType, rangeCd, aoNo, aoPplrName, aoEmailId, aoBldgId, aoBldgDesc
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['CA_reg_number'], row['ENTITY'], row['userId'], row['roleDesc'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                    row['firstName'],  row['midName'], row['lastName'], row['areaCd'], row['areaDesc'], 
                                    row['aoType'], row['rangeCd'], row['aoNo'], row['aoPplrName'], row['aoEmailId'], row['aoBldgId'],
                                    row['aoBldgDesc']
                                ))

                                conn_ef_db.commit()
                                print(f"Juris Diction: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("JURISDCITION DETAILS: CSV CREATED SUCCESSFULLY.")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO juris_diction_info (
                            CA_reg_number, ENTITY, userId, roleDesc, contactFirstName, contactMiddleName, contactLastName, 
                            firstName, midName, lastName, areaCd, areaDesc, 
                            aoType, rangeCd, aoNo, aoPplrName, aoEmailId, aoBldgId, aoBldgDesc
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                        """, (
                            row['CA_reg_number'], row['ENTITY'], row['userId'], row['roleDesc'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                            row['firstName'],  row['midName'], row['lastName'], row['areaCd'], row['areaDesc'], 
                            row['aoType'], row['rangeCd'], row['aoNo'], row['aoPplrName'], row['aoEmailId'], row['aoBldgId'],
                            row['aoBldgDesc']
                        ))

                        conn_ef_db.commit()
                        print(f"Juris Diction: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()
        
        except Exception as e:
            print(f"{username} An error occurred for juris dicition: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_juris_diction
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"JURISDCITION: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"JURISDCITION: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"JURISDCITION: ERROR USERS CSV CREATED.")
        #----------------------------------------JURISDCITION DETAILS END------------------------------------------

        #----------------------------------------SRC. OF INCOME BUSINESS/SALARIED START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_incomesrc = json.dumps({
            "serviceName": "sourceOfIncomeService",
            "userId": response_entity
            })

            headers_incomesrc = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'myPanDetailsService'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            src_response = requests.request("POST", income_source_url, headers = headers_incomesrc, data = payload_incomesrc)
            time.sleep(3)

            src_info = json.loads(src_response.text)
            
            source_type_mapping = {"OTH": "Other", "BUS": "Business / Profession", "HP": "House Property", 
            "AGR": "Agriculture", "SAL": "Salaried/Pensioner"}

            Nature_of_Employment = {"C": "Government", "S": "Pensioners", "P":"PSU", "O":"Other"}
            
            if UserInfo_1.status_code == 200 and src_response.status_code == 200:
                time.sleep(3)
                if src_info:
                    for entry in src_info:
                        if "mbrDetlDtoList" in entry and entry["mbrDetlDtoList"]:
                            time.sleep(3)
                            data_list = []
                            src_type_encrypt = src_info[0]["srcType"]
                            srcType = source_type_mapping.get(src_type_encrypt, "")

                            for mbr_detail in entry["mbrDetlDtoList"]:
                                NatureofEmployment = Nature_of_Employment.get(entry.get("pensionAuthType"))
                                
                                src_business_salaried_info_list = {
                                'userId':basic_info.get('userId', ''),
                                'srcType':srcType,
                                "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                                "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                                "contactLastName":basic_info.get("contactLastName", "")or "Null",
                                "firstName": basic_info.get("firstName", ""),
                                "midName": basic_info.get("midName", ""),
                                "lastName": basic_info.get("lastName", ""),
                                "profileMbrId": mbr_detail.get("profileMbrId", ""),
                                "profileIncmSrcId": mbr_detail.get("profileIncmSrcId", ""),
                                "mbrTanPan": mbr_detail.get("mbrTanPan", ""),
                                "mbrName": mbr_detail.get("mbrName", ""),
                                "NatureofEmployment": NatureofEmployment,
                                }
                                data_list.append(src_business_salaried_info_list)
                        
                            Information = pd.DataFrame(data_list)
                            Information.replace('\x00', '', regex=True, inplace=True)
                                    
                            csv_file_path = src_business_salaried
                            file_exists = os.path.exists(csv_file_path)

                            curr_ef_db.execute("""
                            SELECT EXISTS (
                                SELECT 1
                                FROM information_schema.tables
                                WHERE table_name = 'src_business_salaried'
                            );
                        """)

                        table_exists = curr_ef_db.fetchone()[0]

                        if not table_exists:
                            print("Table does not exist. Creating table...")
                            curr_ef_db.execute("""
                                CREATE TABLE src_business_salaried (
                                    userId text,
                                    srcType text,
                                    contactFirstName text,
                                    contactMiddleName text,
                                    contactLastName text,
                                    firstName text,
                                    midName text,
                                    lastName text,
                                    profileMbrId text,
                                    profileIncmSrcId text,
                                    mbrTanPan text,
                                    mbrName text,
                                    NatureofEmployment text
                                );
                            """)

                            conn_ef_db.commit()
                            print("Table 'src_business_salaried' created successfully.")
                            print()

                            if file_exists:
                                print("File already exists.")

                                if file_exists:
                                    existing_data = pd.read_csv(csv_file_path)
                                    existing_users = existing_data['userId'].tolist()

                                    for index, row in Information.iterrows():
                                        if row['userId'] in existing_users:
                                            print(f"SRC. BUSINESS/SALARIED: {row['userId']} ALREADY EXISTED.")

                                        else:
                                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                                print(f"SRC BUSINESS/SALARIED: {row['userId']} INSERTED SUCCESSFULLY.")

                                            curr_ef_db.execute("""
                                            INSERT INTO src_business_salaried (
                                                userId, srcType, contactFirstName, contactMiddleName, contactLastName, 
                                                firstName, midName, lastName, profileMbrId, profileIncmSrcId, 
                                                mbrTanPan, mbrName, NatureofEmployment
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                            """, (
                                                row['userId'], row['srcType'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                                row['firstName'],  row['midName'], row['lastName'], row['profileMbrId'], row['profileIncmSrcId'],
                                                row['mbrTanPan'], row['mbrName'], row['NatureofEmployment']
                                            ))

                                            conn_ef_db.commit()
                                            print(f"Business/Salaried: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                            print()

                            else:
                                Information.to_csv(csv_file_path, header=True, index=False)
                                print("SRC BUSINESS/SALARIED: CSV CREATED SUCCESSFULLY.")

                                for index, row in Information.iterrows():
                                    curr_ef_db.execute("""
                                    INSERT INTO src_business_salaried (
                                        userId, srcType, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, profileMbrId, profileIncmSrcId, 
                                        mbrTanPan, mbrName, NatureofEmployment
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['userId'], row['srcType'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                        row['firstName'],  row['midName'], row['lastName'], row['profileMbrId'], row['profileIncmSrcId'],
                                        row['mbrTanPan'], row['mbrName'], row['NatureofEmployment']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Business/Salaried: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

        except Exception as e:
            print(f"{username} An error occurred for src. business/salaried: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_src_business_salaried
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"SRC. BUSINESS/SALARIED: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"SRC. BUSINESS/SALARIED: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"SRC. BUSINESS/SALARIED: ERROR USERS CSV CREATED.")   
        #----------------------------------------SRC. OF INCOME BUSINESS/SALARIED END------------------------------------------

        #----------------------------------------SRC. OF INCOME HOUSE PROPERTY START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_incomesrc = json.dumps({
            "serviceName": "sourceOfIncomeService",
            "userId": response_entity
            })

            headers_incomesrc = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            src_response_2 = requests.request("POST", income_source_url, headers = headers_incomesrc, data = payload_incomesrc)
            time.sleep(3)

            src_info_2 = json.loads(src_response_2.text)
            basic_info = json.loads(UserInfo_1.text)

            source_type_mapping = {
                "OTH": "Other",
                "BUS": "Business / Profession",
                "HP": "House Property",
                "AGR": "Agriculture",
                "SAL": "Salaried/Pensioner"
            }

            if UserInfo_1.status_code == 200 and src_response_2.status_code == 200:
                time.sleep(3)
                if src_info_2:
                    for entry_2 in src_info_2:
                        if "propDetlDtoList" in entry_2 and entry_2["propDetlDtoList"]:
                            time.sleep(3)
                            data_list = []
                            src_type_encrypt = src_info_2[0]["srcType"]
                            srcType = source_type_mapping.get(src_type_encrypt, "")

                            for mbr_detail_2 in entry_2["propDetlDtoList"]:

                                stateCd_encrypt = entry_2.get("stateCd")
                                if stateCd_encrypt is not None:
                                    stateCd = state_code.get(stateCd_encrypt, "")
                                else:
                                    stateCd = "Null"

                                countryCd_encrypt = entry_2.get("countryCd")
                                if countryCd_encrypt is not None:
                                    countryCd = country_code.get(countryCd_encrypt, "")
                                else:
                                    countryCd = "Null"

                                src_business_property_info_list = {
                                    'userId':entry_2.get('userId', '') or "Null",
                                    'srcType':srcType,
                                    "contactFirstName":basic_info.get("contactFirstName", "") or "Null",
                                    "contactMiddleName":basic_info.get("contactMiddleName", "") or "Null",
                                    "contactLastName":basic_info.get("contactLastName", "") or "Null",
                                    "firstName": basic_info.get("firstName", "") or "Null",
                                    "midName": basic_info.get("midName", "") or "Null",
                                    "lastName": basic_info.get("lastName", "") or "Null",
                                    'profilePropId':mbr_detail_2.get('profilePropId', '') or "Null",
                                    'profileIncmSrcId':mbr_detail_2.get('profileIncmSrcId', '') or "Null",
                                    'addrLine1Txt':mbr_detail_2.get('addrLine1Txt', '') or "Null",
                                    'pinCd':mbr_detail_2.get('pinCd', '') or "Null",
                                    'stateCd':stateCd,
                                    'countryCd':countryCd,
                                    'ownerPercentage':mbr_detail_2.get('ownPct', '') or "Null",
                                    'noofCoowners':mbr_detail_2.get('noOfCoowner', '') or "Null"
                                }
                                data_list.append(src_business_property_info_list)
                        
                            Information = pd.DataFrame(data_list)
                            Information.replace('\x00', '', regex=True, inplace=True)
                                    
                            csv_file_path = src_house_property
                            file_exists = os.path.exists(csv_file_path)

                            curr_ef_db.execute("""
                                SELECT EXISTS (
                                    SELECT 1
                                    FROM information_schema.tables
                                    WHERE table_name = 'src_house_property'
                                );
                            """)

                            table_exists = curr_ef_db.fetchone()[0]

                            if not table_exists:
                                print("Table does not exist. Creating table...")
                                curr_ef_db.execute("""
                                    CREATE TABLE src_house_property (
                                        userId text,
                                        srcType text,
                                        contactFirstName text,
                                        contactMiddleName text,
                                        contactLastName text,
                                        firstName text,
                                        midName text,
                                        lastName text,
                                        profilePropId text,
                                        profileIncmSrcId text,
                                        addrLine1Txt text,
                                        pinCd text,
                                        stateCd text,
                                        countryCd text,
                                        ownerPercentage text,
                                        noofCoowners text
                                    );
                                """)

                                conn_ef_db.commit()
                                print("Table 'src_house_property' created successfully.")
                                print()

                            if file_exists:
                                existing_data = pd.read_csv(csv_file_path)
                                existing_users = existing_data['userId'].tolist()

                                for index, row in Information.iterrows():
                                    if row['userId'] in existing_users:
                                        print(f"SRC. HOUSE PROPERTY: {row['userId']} ALREADY EXISTED.")

                                    else:
                                        with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                            Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                            print(f"SRC HOUSE PROPERTY: {row['userId']} INSERTED SUCCESSFULLY.")

                                        curr_ef_db.execute("""
                                        INSERT INTO src_house_property (
                                            userId, srcType, contactFirstName, contactMiddleName, contactLastName, 
                                            firstName, midName, lastName, profilePropId, profileIncmSrcId, 
                                            addrLine1Txt, pinCd, stateCd, countryCd, ownerPercentage, noofCoowners
                                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                        """, (
                                            row['userId'], row['srcType'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                            row['firstName'],  row['midName'], row['lastName'], row['profilePropId'], row['profileIncmSrcId'],
                                            row['addrLine1Txt'], row['pinCd'], row['stateCd'], row['countryCd'], row['ownerPercentage'], row['noofCoowners']
                                        ))

                                        conn_ef_db.commit()
                                        print(f"Src House Property: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                        print()

                            else:
                                Information.to_csv(csv_file_path, header=True, index=False)
                                print("SRC HOUSE PROPERTY: CSV CREATED SUCCESSFULLY.")

                                for index, row in Information.iterrows():
                                    curr_ef_db.execute("""
                                    INSERT INTO src_house_property (
                                        userId, srcType, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, profilePropId, profileIncmSrcId, 
                                        addrLine1Txt, pinCd, stateCd, countryCd, ownerPercentage, noofCoowners
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['userId'], row['srcType'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                        row['firstName'],  row['midName'], row['lastName'], row['profilePropId'], row['profileIncmSrcId'],
                                        row['addrLine1Txt'], row['pinCd'], row['stateCd'], row['countryCd'], row['ownerPercentage'], row['noofCoowners']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Src House Property: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()
        
        except Exception as e:
            print(f"{username} An error occurred for src. house property: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_src_house_property
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"SRC. HOUSE PROPERTY: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"SRC. HOUSE PROPERTY: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"SRC. HOUSE PROPERTY: ERROR USERS CSV CREATED.")
        #----------------------------------------SRC. OF HOUSE PROPERTY END------------------------------------------

        #----------------------------------------AUTHORISED SIGNATURE START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_authsign = json.dumps({
                "header": {
                    "formName": "FO-049-MYPFL"
                }, 
                "serviceName": "profileauthPrsnDetl"
            })

            headers_authsign = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn':"profileauthPrsnDetl"
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            AuthSign = requests.request("POST", auth_sign_url, headers = headers_authsign, data = payload_authsign)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            authsign = json.loads(AuthSign.text)
            time.sleep(3)
            
            if UserInfo_1.status_code == 200 and AuthSign.status_code == 200:
                auth_sign_list = [{
                    "userId": basic_info.get("userId", ""),
                    "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                    "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                    "contactLastName":basic_info.get("contactLastName", "")or "Null",
                    "firstName": basic_info.get("firstName", ""),
                    "midName": basic_info.get("midName", ""),
                    "lastName": basic_info.get("lastName", ""),
                    "roleDesc": basic_info.get("roleDesc", ""),
                    "authRepPan":authsign.get("authRepPan", "")or "Not Define",
                    "authRepFirstNm":authsign.get("authRepFirstNm", "") or "Not Define",
                    "authRepMidNm":authsign.get("authRepMidNm", "")or "Not Define",
                    "authRepLastNm":authsign.get("authRepLastNm", "")or "Not Define",
                    "periodTo":authsign.get("periodTo", "")or "Not Define",
                    "periodFrom":authsign.get("periodFrom", "")or "Not Define",
                    "dscFlag":authsign.get("dscFlag", "")or "Not Define",
                    "dscExpDt":authsign.get("dscExpDt", "")or "Not Define",
                    "taskAssigned":authsign.get("taskAssigned", "")or "Not Define",
                    "reason":authsign.get("taskAssigned", "")or "Not Define"
                }]

                Information = pd.DataFrame(auth_sign_list)
                Information.replace('\x00', '', regex=True, inplace=True)
                        
                file_exists = os.path.exists(base_dir)
                
                csv_file_path = authorised_signature
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'auth_sign'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE auth_sign (
                            userId text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            firstName text,
                            midName text,
                            lastName text,
                            roleDesc text,
                            authRepPan text,
                            authRepFirstNm text,
                            authRepMidNm text,
                            authRepLastNm text,
                            periodTo text,
                            periodFrom text,
                            dscFlag text,
                            dscExpDt text,
                            taskAssigned" text,
                            reason text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'auth_sign' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"AUTHORISED SIGNATURE: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"AUTHORISED SIGNATURE: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO auth_sign (
                                    userId, contactFirstName, contactMiddleName, contactLastName, 
                                    firstName, midName, lastName, roleDesc, authRepPan, authRepFirstNm, 
                                    authRepMidNm, authRepLastNm, periodTo, periodFrom, dscFlag, dscExpDt, 
                                    taskAssigned, reason
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['userId'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                    row['firstName'],  row['midName'], row['lastName'], row['roleDesc'], row['authRepPan'], row['authRepFirstNm'], 
                                    row['authRepMidNm'], row['authRepLastNm'], row['periodTo'], row['periodFrom'], row['dscFlag'], row['dscExpDt'],
                                    row['taskAssigned'], row['reason']
                                ))

                                conn_ef_db.commit()
                                print(f"Authorised Signature: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("AUTHORISED SIGNATURE: CSV CREATED SUCCESSFULLY.")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO auth_sign (
                            userId, contactFirstName, contactMiddleName, contactLastName, 
                            firstName, midName, lastName, roleDesc, authRepPan, authRepFirstNm, 
                            authRepMidNm, authRepLastNm, periodTo, periodFrom, dscFlag, dscExpDt, 
                            taskAssigned, reason
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                        """, (
                            row['userId'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                            row['firstName'],  row['midName'], row['lastName'], row['roleDesc'], row['authRepPan'], row['authRepFirstNm'], 
                            row['authRepMidNm'], row['authRepLastNm'], row['periodTo'], row['periodFrom'], row['dscFlag'], row['dscExpDt'],
                            row['taskAssigned'], row['reason']
                        ))

                        conn_ef_db.commit()
                        print(f"Authorised Signature: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()
            
        except Exception as e:
            print(f"{username} An error occurred for authorised signature: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_authorised_signature
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"AUTHORISED SIGNATURE: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"AUTHORISED SIGNATURE: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"AUTHORISED SIGNATURE: ERROR USERS CSV CREATED.")
        #----------------------------------------AUTHORISED SIGNATURE END------------------------------------------

        #----------------------------------------REPRESENTATIVE ASSES START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_Repassess = json.dumps({
            "header": {
                "formName": "FO-049-MYPFL"
            },
            "serviceName": "profileRepAssesse"
            })

            headers_Repassess = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn':"profileauthPrsnDetl"
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            RepAssess = requests.request("POST", rep_assess_url, headers = headers_Repassess, data = payload_Repassess)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            repassess = json.loads(RepAssess.text)
            time.sleep(3)
            
            if UserInfo_1.status_code == 200 and RepAssess.status_code == 200:
                rep_asses_list = [{
                    "userId": basic_info.get("userId", ""),
                    "roleDesc": basic_info.get("roleDesc", ""),
                    "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                    "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                    "contactLastName":basic_info.get("contactLastName", "")or "Null",
                    "firstName": basic_info.get("firstName", ""),
                    "midName": basic_info.get("midName", ""),
                    "lastName": basic_info.get("lastName", ""),
                    "authRepPan":repassess.get("authRepPan", "")or "Null",
                    "authRepFirstNm":repassess.get("authRepFirstNm", "") or "Null",
                    "authRepMidNm":repassess.get("authRepMidNm", "")or "Null",
                    "authRepLastNm":repassess.get("authRepLastNm", "")or "Null",
                    "periodTo":repassess.get("periodTo", "")or "Null",
                    "periodFrom":repassess.get("periodFrom", "")or "Null",
                    "dscFlag":repassess.get("dscFlag", "")or "Null",
                    "dscExpDt":repassess.get("dscExpDt", "")or "Null",
                    "taskAssigned":repassess.get("taskAssigned", "")or "Null",
                    "reason":repassess.get("taskAssigned", "")or "Null"
                }]

                Information = pd.DataFrame(rep_asses_list)
                Information.replace('\x00', '', regex=True, inplace=True)
                        
                csv_file_path = representative_asseses
                file_exists = os.path.exists(csv_file_path)

                curr_ef_db.execute("""
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.tables
                        WHERE table_name = 'representative_assessee'
                    );
                """)

                table_exists = curr_ef_db.fetchone()[0]

                if not table_exists:
                    print("Table does not exist. Creating table...")
                    curr_ef_db.execute("""
                        CREATE TABLE representative_assessee (
                            userId text,
                            roleDesc text,
                            contactFirstName text,
                            contactMiddleName text,
                            contactLastName text,
                            firstName text,
                            midName text,
                            lastName text,
                            authRepPan text,
                            authRepFirstNm text,
                            authRepMidNm text,
                            authRepLastNm text,
                            periodTo text,
                            periodFrom text,
                            dscFlag text,
                            dscExpDt text,
                            taskAssigned text,
                            reason text
                        );
                    """)

                    conn_ef_db.commit()
                    print("Table 'representative_assessee' created successfully.")
                    print()

                if file_exists:
                    print("File already exists.")

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"REPRESENTATIVE ASSES: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"REPRESENTATIVE ASSESES: {row['userId']} INSERTED SUCCESSFULLY.")

                                curr_ef_db.execute("""
                                INSERT INTO representative_assessee (
                                    userId, roleDesc, contactFirstName, contactMiddleName, contactLastName, 
                                    firstName, midName, lastName, authRepPan, authRepFirstNm, 
                                    authRepMidNm, authRepLastNm, periodTo, periodFrom, 
                                    dscFlag, dscExpDt, taskAssigned, reason
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['userId'], row['roleDesc'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                                    row['firstName'],  row['midName'], row['lastName'], row['authRepPan'], row['authRepFirstNm'], 
                                    row['authRepMidNm'], row['authRepLastNm'], row['periodTo'], row['periodFrom'], row['dscFlag'], row['dscExpDt'],
                                    row['taskAssigned'], row['reason']
                                ))

                                conn_ef_db.commit()
                                print(f"Representative Assessee: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print("REPRESENTATIVE ASSESES: CSV CREATED SUCCESSFULLY.")

                    for index, row in Information.iterrows():
                        curr_ef_db.execute("""
                        INSERT INTO representative_assessee (
                            userId, roleDesc, contactFirstName, contactMiddleName, contactLastName, 
                            firstName, midName, lastName, authRepPan, authRepFirstNm, 
                            authRepMidNm, authRepLastNm, periodTo, periodFrom, 
                            dscFlag, dscExpDt, taskAssigned, reason
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                        """, (
                            row['userId'], row['roleDesc'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'], 
                            row['firstName'],  row['midName'], row['lastName'], row['authRepPan'], row['authRepFirstNm'], 
                            row['authRepMidNm'], row['authRepLastNm'], row['periodTo'], row['periodFrom'], row['dscFlag'], row['dscExpDt'],
                            row['taskAssigned'], row['reason']
                        ))

                        conn_ef_db.commit()
                        print(f"Representative Assessee: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                        print()
            
        except Exception as e:
            print(f"{username} An error occurred for representative asses: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_representative_asses
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"REPRESENTATIVE ASSES: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"REPRESENTATIVE ASSES: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"REPRESENTATIVE ASSES: ERROR USERS CSV CREATED.")
        #----------------------------------------REPRESENTATIVE ASSES END------------------------------------------

        #----------------------------------------DEMAT ACCOUNT START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_DematAcc = json.dumps({
            "entityNum": response_entity,
            "serviceName": "myDematAccountDetailsServiceImpl",
            "header": {
                "formName": "FO-055-PDACC"
            }
            })

            headers_DematAcc = {
            'Accept': 'application/json, text/plain, */*',
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            DematAcc = requests.request("POST", demat_acc_url, headers = headers_DematAcc, data = payload_DematAcc)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            dematacc = json.loads(DematAcc.text)

            #----------------------------------------ACTIVE DEMAT ACCOUNT START------------------------------------------
            try:
                time.sleep(3)
                if dematacc.get('activeDematAccList', []):    
                    active_acc = dematacc.get('activeDematAccList')
                    active_demat_data_list = []
                    for entry in active_acc:
                        encrypt_acc_no = entry.get('dematAccountNumber')
                        acc_no = base64.b64decode(encrypt_acc_no).decode('utf-8')

                        encrypt_mobile_no = entry.get('mobileNo')
                        mobile_no = base64.b64decode(encrypt_mobile_no).decode('utf-8')

                        encrypt_email = entry.get('emailId')
                        email = base64.b64decode(encrypt_email).decode('utf-8')

                        active_demat_acc_list = {
                        'userId':basic_info.get('userId'),
                        'accountNo':acc_no,
                        'depType':entry.get('depType', ''),
                        'MobileNo':mobile_no,
                        'EmailId':email,
                        'nameAsPerDemat':entry.get('nameAsPerDemat', ''),
                        'nameVerFlag':entry.get('nameVerFlag', ''),
                        'mobileVerFlag':entry.get('mobileVerFlag', ''),
                        'emailVerFlag':entry.get('emailVerFlag', ''),
                        'transactionNo':entry.get('transactionNo', ''),
                        'panLinkFlag':entry.get('panLinkFlag', ''),
                        'dematPanLinkingId':entry.get('dematPanLinkingId', ''),
                        }
                        active_demat_data_list.append(active_demat_acc_list)
                
                    Information = pd.DataFrame(active_demat_data_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                            
                    csv_file_path = active_demat_account
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'active_demat_account'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE active_demat_account (
                                userId text,
                                accountNo text,
                                depType text,
                                MobileNo text,
                                EmailId text,
                                nameAsPerDemat text,
                                nameVerFlag text,
                                mobileVerFlag text,
                                emailVerFlag text,
                                transactionNo text,
                                panLinkFlag text,
                                dematPanLinkingId text,
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'active_demat_account' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['userId'].tolist()

                            for index, row in Information.iterrows():
                                if row['userId'] in existing_users:
                                    print(f"ACTIVE DEMAT ACCOUNT: {row['userId']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"ACTIVE DEMAT ACCOUNT: {row['userId']} INSERTED SUCCESSFULLY.")

                                    curr_ef_db.execute("""
                                    INSERT INTO active_demat_account (
                                        userId, accountNo, depType, MobileNo, EmailId, nameAsPerDemat, 
                                        nameVerFlag, mobileVerFlag, emailVerFlag, transactionNo, 
                                        panLinkFlag, dematPanLinkingId
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['userId'], row['accountNo'], row['depType'], row['MobileNo'], row['EmailId'], 
                                        row['nameAsPerDemat'], row['nameVerFlag'], row['mobileVerFlag'], row['emailVerFlag'], 
                                        row['transactionNo'], row['panLinkFlag'], row['dematPanLinkingId']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"Active Demat Account: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("ACTIVE DEMAT ACCOUNT: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO active_demat_account (
                                userId, accountNo, depType, MobileNo, EmailId, nameAsPerDemat, 
                                nameVerFlag, mobileVerFlag, emailVerFlag, transactionNo, 
                                panLinkFlag, dematPanLinkingId
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['userId'], row['accountNo'], row['depType'], row['MobileNo'], row['EmailId'], 
                                row['nameAsPerDemat'], row['nameVerFlag'], row['mobileVerFlag'], row['emailVerFlag'], 
                                row['transactionNo'], row['panLinkFlag'], row['dematPanLinkingId']
                            ))

                            conn_ef_db.commit()
                            print(f"Active Demat Account: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()

            except Exception as e:
                print(f"{username} An error occurred for active demat account: {e}")
                name = username
                passwrd = password

                user_basic_info_list = [{ 
                    "ID":name,
                    "Password":passwrd
                }]

                Information = pd.DataFrame(user_basic_info_list)
                Information.replace('\x00', '', regex=True, inplace=True)

                csv_file_path = Error_active_demat_account
                file_exists = os.path.exists(csv_file_path)
            
                if file_exists:
                    print("File already exists.")
                                
                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['ID'].tolist()

                        for index, row in Information.iterrows():
                            if row['ID'] in existing_users:
                                print(f"ACTIVE DEMAT ACCOUNT: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"ACTIVE DEMAT ACCOUNT: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
            
                else:
                    Information.to_csv(csv_file_path, header=True, index=False)
                    print(f"ACTIVE DEMAT ACCOUNT: ERROR USERS CSV CREATED.")
            #----------------------------------------ACTIVE DEMAT ACCOUNT END------------------------------------------
                        
        except Exception as e:
            print(f"{username} An error occurred for demat account: {e}")
            name = username
            passwrd = password

            user_basic_info_list = [{ 
            "ID":name,
            "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_demat_account
            file_exists = os.path.exists(csv_file_path)
            
            if file_exists:
                print("File already exists.")
                                
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"DEMAT ACCOUNT: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"DEMAT ACCOUNT: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
            
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"DEMAT ACCOUNT: ERROR USERS CSV CREATED.")
        #----------------------------------------DEMAT ACCOUNT END------------------------------------------

        #----------------------------------------FYA START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_eproceedingCount = json.dumps({
            "serviceName": "getEProceedingsDetail",
            "pan": response_entity,
            "prcdngStatusFlag": "FYA",
            "header": {
                "formName": "FO-041_PCDNG"
            }
            })

            headers_eproceedingCount = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            }

            
            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            Response_FYA = requests.request("POST", fya_count_url, headers = headers_eproceedingCount, data = payload_eproceedingCount)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            response_fya = json.loads(Response_FYA.text)
            time.sleep(3)

            if UserInfo_1.status_code == 200 and Response_FYA.status_code == 200:
                # ------------------------------FYA NOTICE COUNT START------------------------------
                try:
                    time.sleep(3)
                    fay_count_list = [{
                    "userId": basic_info.get("userId", ""),
                    "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                    "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                    "contactLastName":basic_info.get("contactLastName", "")or "Null",
                    "firstName": basic_info.get("firstName", "")or "Null",
                    "midName": basic_info.get("midName", "")or "Null",
                    "lastName": basic_info.get("lastName", "")or "Null",
                    "roleDesc": basic_info.get("roleDesc", "")or "Null",
                    "eProceedingsInfoCount" : response_fya.get("eProceedingsInfoCount", "")or "Null",
                    "eProceedingsForActionCount" : response_fya.get("eProceedingsForActionCount", "")or "Null"
                    }]
                
                    Information = pd.DataFrame(fay_count_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                                    
                    csv_file_path = fya_notice_count
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'fya_count'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE fya_count (
                                userId text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                roleDesc text,
                                eProceedingsInfoCount text,
                                eProceedingsForActionCount text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'fya_count' created successfully.")
                        print()

                    if file_exists:
                        existing_data = pd.read_csv(csv_file_path)
                        existing_users = existing_data['userId'].tolist()

                        for index, row in Information.iterrows():
                            if row['userId'] in existing_users:
                                print(f"FYA NOTICE COUNT: {row['userId']} ALREADY EXISTED.")

                            else:
                                with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                    Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                    print(f"FYA NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY.")
                                
                                curr_ef_db.execute("""
                                    INSERT INTO fya_count (
                                        userId, contactFirstName, contactMiddleName, contactLastName, firstName, midName, lastName, roleDesc, eProceedingsInfoCount, eProceedingsForActionCount
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['userId'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'],
                                    row['firstName'], row['midName'], row['lastName'], row['roleDesc'],
                                    row['eProceedingsInfoCount'], row['eProceedingsForActionCount']
                                ))

                                conn_ef_db.commit()
                                print(f"FYA NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("FYA NOTICE COUNT: CSV CREATED SUCCESSFULLY.")
                        print()

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                                INSERT INTO fya_count (
                                    userId, contactFirstName, contactMiddleName, contactLastName, firstName, midName, lastName, roleDesc, eProceedingsInfoCount, eProceedingsForActionCount
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['userId'], row['contactFirstName'], row['contactMiddleName'], row['contactLastName'],
                                row['firstName'], row['midName'], row['lastName'], row['roleDesc'],
                                row['eProceedingsInfoCount'], row['eProceedingsForActionCount']
                            ))

                            conn_ef_db.commit()
                            print(f"FYA NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()

                except Exception as e:
                    print(f"{username} An error occurred for FYA count notice: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fya_notice_count
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYA NOTICE COUNT: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA NOTICE COUNT: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA NOTICE COUNT: ERROR USERS CSV CREATED.")
                #------------------------------FYA NOTICE COUNT END------------------------------

                #------------------------------FYA NOTICE DESCRIPTION START------------------------------
                try:
                    time.sleep(3)
                    if response_fya.get('eproceedingRequests', []):
                        eprorequest = response_fya.get('eproceedingRequests', [])
                        data_list = []

                        for entry in eprorequest:

                            proceedingLimitationDate_encrypt = entry.get("proceedingLimitationDate", "")
                            if proceedingLimitationDate_encrypt:
                                proceedingLimitationDate_Date = datetime.fromtimestamp(proceedingLimitationDate_encrypt / 1000)
                                proceedingLimitationDate = proceedingLimitationDate_Date.strftime('%d-%m-%Y')
                                proceedingLimitationDateTime = proceedingLimitationDate_Date.time()
                            else:
                                proceedingLimitationDate = "Null"
                                proceedingLimitationDateTime = "Null"

                            responseDate_encrypt = entry.get("responseDate", "")
                            if responseDate_encrypt:
                                responseDate_Date = datetime.fromtimestamp(responseDate_encrypt / 1000)
                                responseDate = responseDate_Date.strftime('%d-%m-%Y')
                                responseDateTime = responseDate_Date.time()
                            else:
                                responseDate = "Null"
                                responseDateTime = "Null"

                            acknowledgementNo_encrypt =  entry.get("acknowledgementNo", "")
                            if acknowledgementNo_encrypt:
                                acknowledgementNo = f"'{acknowledgementNo_encrypt}"
                            else:
                                acknowledgementNo = "Null"


                            issuedOn_encrypt = entry.get("issuedOn", "")
                            if issuedOn_encrypt:
                                issuedOn_Date = datetime.fromtimestamp(issuedOn_encrypt / 1000)
                                issuedOn = issuedOn_Date.strftime('%d-%m-%Y')
                                issudedOnTime = issuedOn_Date.time()
                            else:
                                issuedOn = "Null"
                                issuedOnTime = "Null"
                            
            
                            servedOn_encrypt = entry.get("servedOn", "")
                            if servedOn_encrypt:
                                servedOn_Date = datetime.fromtimestamp(servedOn_encrypt / 1000)
                                servedOn = servedOn_Date.strftime('%d-%m-%Y')
                                servedOnTime = servedOn_Date.time()
                            else:
                                servedOn = "Null"
                                servedOnTime = "Null"
            
                            responseDueDate_encrypt = entry.get("responseDueDate", "")
                            if responseDueDate_encrypt:
                                responseDueDate_Date = datetime.fromtimestamp(responseDueDate_encrypt / 1000)
                                responseDueDate = responseDueDate_Date.strftime('%d-%m-%Y')
                                responseDueDateTime = responseDueDate_Date.time()
                            else:
                                responseDueDate = "Null"
                                responseDueDateTime = "Null"
                            
                            lastResponseSubmittedOn_encrypt = entry.get("lastResponseSubmittedOn", "")
                            if lastResponseSubmittedOn_encrypt:
                                lastResponseSubmittedOn_Date = datetime.fromtimestamp(lastResponseSubmittedOn_encrypt / 1000)
                                lastResponseSubmittedOn = lastResponseSubmittedOn_Date.strftime('%d-%m-%Y')
                                lastResponseSubmittedOnTime = lastResponseSubmittedOn_Date.time()
                            else:
                                lastResponseSubmittedOn = "Null"
                                lastResponseSubmittedOnTime = "Null"

                            proceedingClosureDate_encrypt = entry.get("proceedingClosureDate", "")
                            if proceedingClosureDate_encrypt:
                                proceedingClosureDate_Date = datetime.fromtimestamp(proceedingClosureDate_encrypt / 1000)
                                proceedingClosureDate = proceedingClosureDate_Date.strftime('%d-%m-%Y')
                                proceedingClosureDateTime = proceedingClosureDate_Date.time()
                            else:
                                proceedingClosureDate = "Null"
                                proceedingClosureDateTime = "Null"

                        
                            fya_notice_description_list = {
                            'proceedingReqId': entry.get("proceedingReqId", ""),
                            'pan': entry.get("pan", ""),
                            'nameOfAssesse': entry.get("nameOfAssesse", ""),
                            'proceedingName':entry.get("proceedingName", ""), 
                            'itrType':entry.get("itrType", ""),
                            'assessmentYear':entry.get("assessmentYear", ""),
                            'financialYr':entry.get("financialYr", ""),
                            'proceedingLimitationDate':proceedingLimitationDate,
                            'proceedingLimitationDateTime':proceedingLimitationDateTime,
                            'noticeName':entry.get("noticeName", ""),
                            'responseDate':responseDate,
                            'responseDateTime':responseDateTime,
                            'acknowledgementNo':acknowledgementNo,
                            'viewNoticeCount':entry.get("viewNoticeCount", ""),
                            'proceedingType':entry.get("proceedingType", ""),
                            'issuedOn':issuedOn,
                            'issuedOnTime':issuedOnTime,
                            "servedOn" : servedOn,
                            "servedOnTime":servedOnTime,
                            "responseDueDate" : responseDueDate,
                            "responseDueDateTime":responseDueDateTime,
                            "lastResponseSubmittedOn" : lastResponseSubmittedOn,
                            "lastResponseSubmittedOnTime":lastResponseSubmittedOnTime,
                            'responseViewedByAoOn':entry.get("responseViewedByAoOn", ""),
                            'proceedingClosureDate':proceedingClosureDate,
                            'proceedingClosureDateTime':proceedingClosureDateTime,
                            'proceedingClosureOrder':entry.get("proceedingClosureOrder", "") or "Null",
                            'proceedingStatus':entry.get("proceedingStatus", ""),
                            'respStatus':entry.get("respStatus", ""),
                            'respId':entry.get("respId", ""),
                            'commType':entry.get("commType", ""),
                            'readFlag':entry.get("readFlag", ""),
                            'facelessFlag':entry.get("facelessFlag", ""),
                            'returnEverified':entry.get("returnEverified", "")or "Null",
                            'discardAllowed':entry.get("discardAllowed", "")or "Null",
                            'new':entry.get("new", ""),
                            'isFileAppeal':entry.get("isFileAppeal", ""),
                            'isRectification':entry.get("isRectification", ""),
                            }
                            data_list.append(fya_notice_description_list)

                        Information = pd.DataFrame(data_list)
                        Information.replace('\x00', '', regex=True, inplace=True)

                        csv_file_path = fya_notice_description
                        file_exists = os.path.exists(csv_file_path)

                        curr_ef_db.execute("""
                            SELECT EXISTS (
                                SELECT 1
                                FROM information_schema.tables
                                WHERE table_name = 'fya_notice_description'
                            );
                        """)

                        table_exists = curr_ef_db.fetchone()[0]

                        if not table_exists:
                            print("Table does not exist. Creating table...")
                            curr_ef_db.execute("""
                                CREATE TABLE fya_notice_description (
                                    proceedingReqId text,
                                    pan text,
                                    nameOfAssesse text,
                                    proceedingName text,
                                    itrType text,
                                    assessmentYear text,
                                    financialYr text,
                                    proceedingLimitationDate text,
                                    proceedingLimitationDateTime text,
                                    noticeName text,
                                    responseDate text,
                                    responseDateTime text,
                                    acknowledgementNo text,
                                    viewNoticeCount text,
                                    proceedingType text,
                                    issuedOn text,
                                    issuedOnTime text,
                                    servedOn text,
                                    servedOnTime text,
                                    responseDueDate text,
                                    responseDueDateTime text,
                                    lastResponseSubmittedOn text,
                                    lastResponseSubmittedOnTime text,
                                    responseViewedByAoOn text,
                                    proceedingClosureDate text,
                                    proceedingClosureDateTime text,
                                    proceedingClosureOrder text,
                                    proceedingStatus text,
                                    respStatus text,
                                    respId text,
                                    commType text,
                                    readFlag text,
                                    facelessFlag text,
                                    returnEverified text,
                                    discardAllowed text,
                                    new text,
                                    isFileAppeal text,
                                    isRectification text
                                );
                            """)

                            conn_ef_db.commit()
                            print("Table 'fya_notice_description' created successfully.")
                            print()

                        if file_exists:
                            print("File already exists.")

                            if file_exists:
                                existing_data = pd.read_csv(csv_file_path)
                                existing_users = existing_data['proceedingReqId'].tolist()

                                for index, row in Information.iterrows():
                                    if row['proceedingReqId'] in existing_users:
                                        print(f"FYA NOTICE DESCRIPTION: {row['pan']} {row['proceedingReqId']} ALREADY EXISTED.")

                                    else:
                                        with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                            Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                            print(f"FYA NOTICE DESCRIPTION: {row['pan']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")

                                        curr_ef_db.execute("""
                                            INSERT INTO fya_notice_description (
                                                proceedingReqId, pan, nameOfAssesse, proceedingName, itrType, assessmentYear, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, 
                                                noticeName, responseDate, responseDateTime, acknowledgementNo, viewNoticeCount, proceedingType, issuedOn, issuedOnTime, servedOn, servedOnTime, responseDueDate,
                                                responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, proceedingClosureDate, proceedingClosureDateTime, proceedingClosureOrder,
                                                proceedingStatus, respStatus, respId, commType, readFlag, facelessFlag, returnEverified, discardAllowed, new, isFileAppeal, isRectification
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                        """, (
                                            row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['proceedingName'], row['itrType'], row['assessmentYear'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'],
                                            row['noticeName'], row['responseDate'], row['responseDateTime'], row['acknowledgementNo'],  row['viewNoticeCount'], row['proceedingType'], row['issuedOn'], row['issuedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                            row['responseDueDateTime'], row['lastResponseSubmittedOn'],row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['proceedingClosureDate'],row['proceedingClosureDateTime'], row['proceedingClosureOrder'],
                                            row['proceedingStatus'], row['respStatus'], row['respId'], row['commType'], row['readFlag'],row['facelessFlag'], row['returnEverified'], row['discardAllowed'], row['new'],row['isFileAppeal'], row['isRectification']
                                        ))

                                        conn_ef_db.commit()
                                        print(f"FYA Notice Description: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                        print()

                        else:
                            Information.to_csv(csv_file_path, header=True, index=False)
                            print("FYA NOTICE DESCRIPTION: CSV CREATED SUCCESSFULLY.")

                            for index, row in Information.iterrows():
                                curr_ef_db.execute("""
                                    INSERT INTO fya_notice_description (
                                        proceedingReqId, pan, nameOfAssesse, proceedingName, itrType, assessmentYear, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, 
                                        noticeName, responseDate, responseDateTime, acknowledgementNo, viewNoticeCount, proceedingType, issuedOn, issuedOnTime, servedOn, servedOnTime, responseDueDate,
                                        responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, proceedingClosureDate, proceedingClosureDateTime, proceedingClosureOrder,
                                        proceedingStatus, respStatus, respId, commType, readFlag, facelessFlag, returnEverified, discardAllowed, new, isFileAppeal, isRectification
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                """, (
                                    row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['proceedingName'], row['itrType'], row['assessmentYear'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'],
                                    row['noticeName'], row['responseDate'], row['responseDateTime'], row['acknowledgementNo'],  row['viewNoticeCount'], row['proceedingType'], row['issuedOn'], row['issuedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                    row['responseDueDateTime'], row['lastResponseSubmittedOn'],row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['proceedingClosureDate'],row['proceedingClosureDateTime'], row['proceedingClosureOrder'],
                                    row['proceedingStatus'], row['respStatus'], row['respId'], row['commType'], row['readFlag'],row['facelessFlag'], row['returnEverified'], row['discardAllowed'], row['new'],row['isFileAppeal'], row['isRectification']
                                ))

                                conn_ef_db.commit()
                                print(f"FYA Notice Description: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                print()
                
                except Exception as e:
                    print(f"{username} An error occurred for FYA notice description: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fya_notice_description
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYA NOTICE DESCRIPTION: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA NOTICE DESCRIPTION: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA NOTICE DESCRIPTION: ERROR USERS CSV CREATED.")
                #------------------------------FYA NOTICE DESCRIPTION END------------------------------

                #------------------------------FYA NOTICE ALL NOTICES START------------------------------
                try:
                    time.sleep(3)
                    if response_fya.get('eproceedingRequests', []):
                        for proceeding_request in response_fya.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)
                        
                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }

                            response_view_notices = requests.request("POST", fya_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)
                            
                            view_notice = json.loads(response_view_notices.text)

                            num_entries = len(view_notice)
                            
                            data_list = []
                            for entry in view_notice:

                                documentIdentificationNumber_encrypt =  entry.get("documentIdentificationNumber", "")
                                if documentIdentificationNumber_encrypt is not None and documentIdentificationNumber_encrypt != "":
                                    if isinstance(documentIdentificationNumber_encrypt, str):
                                        documentIdentificationNumber = documentIdentificationNumber_encrypt
                                    else:
                                        documentIdentificationNumber = f"'{documentIdentificationNumber_encrypt}"
                                else:
                                    documentIdentificationNumber = "Null"

                                proceedingLimitationDate_encrypt = entry.get("proceedingLimitationDate", "")
                                if proceedingLimitationDate_encrypt is not None and proceedingLimitationDate_encrypt in entry:
                                    proceedingLmtDate = datetime.fromtimestamp(proceedingLimitationDate_encrypt / 1000)
                                    proceedingLimitationDate = proceedingLmtDate.strftime('%d-%m-%Y')
                                    proceedingLimitationDateTime = proceedingLmtDate.time()
                                else:
                                    proceedingLimitationDate = "Null"
                                    proceedingLimitationDateTime = "Null"

                                issuedOn_encrypt = entry.get("issuedOn", "")
                                if issuedOn_encrypt is not None:
                                    issuedDate = datetime.fromtimestamp(issuedOn_encrypt / 1000)
                                    issuedOn = issuedDate.strftime('%d-%m-%Y')
                                    issudedOnTime = issuedDate.time()
                                else:
                                    issuedOn = "Null"
                                    issudedOnTime = "Null"

                                servedOn_encrypt = entry.get("servedOn", "")
                                if servedOn_encrypt is not None:
                                    servedDate = datetime.fromtimestamp(servedOn_encrypt / 1000)
                                    servedOn = servedDate.strftime('%d-%m-%Y')
                                    servedOnTime = servedDate.time()
                                else:
                                    servedOn = "Null"
                                    servedOnTime = "Null"

                                responseDueDate_encrypt = entry.get("responseDueDate", "")
                                if responseDueDate_encrypt is not None:
                                    responseDate = datetime.fromtimestamp(responseDueDate_encrypt / 1000)
                                    responseDueDate = responseDate.strftime('%d-%m-%Y')
                                    responseDueDateTime = responseDate.time()
                                else:
                                    responseDueDate = "Null"
                                    responseDueDateTime = "Null"

                                lastResponseSubmittedOn_encrypt = entry.get("lastResponseSubmittedOn", "")
                                if lastResponseSubmittedOn_encrypt is not None:
                                    lastResponseDate = datetime.fromtimestamp(lastResponseSubmittedOn_encrypt / 1000)
                                    lastResponseSubmittedOn = lastResponseDate.strftime('%d-%m-%Y')
                                    lastResponseSubmittedOnTime = lastResponseDate.time()
                                else:
                                    lastResponseSubmittedOn = "Null"
                                    lastResponseSubmittedOnTime = "Null"
                                
                                fya_all_notices = {
                                'proceedingReqId': entry.get("proceedingReqId", ""),
                                'pan': entry.get("pan", ""),
                                'nameOfAssesse': entry.get("nameOfAssesse", ""),
                                'headerSeqNo': entry.get("headerSeqNo", ""),
                                'proceedingName':entry.get("proceedingName", ""),
                                'financialYr':entry.get("financialYr", ""),
                                'proceedingLimitationDate':proceedingLimitationDate,
                                'proceedingLimitationDateTime':proceedingLimitationDateTime,
                                'proceedingType':entry.get("proceedingType", ""),
                                'documentIdentificationNumber':documentIdentificationNumber,
                                'ay':entry.get("ay", ""),
                                'noticeSection':entry.get("noticeSection", ""),
                                'description':entry.get("description", ""),
                                "issuedOn" : issuedOn,
                                "issudedOnTime": issudedOnTime,
                                "servedOn" : servedOn,
                                "servedOnTime":servedOnTime,
                                "responseDueDate" : responseDueDate,
                                "responseDueDateTime":responseDueDateTime,
                                "lastResponseSubmittedOn" : lastResponseSubmittedOn,
                                "lastResponseSubmittedOnTime":lastResponseSubmittedOnTime,
                                'responseViewedByAoOn':entry.get("responseViewedByAoOn", ""),
                                'documentReferenceId':entry.get("documentReferenceId", ""),
                                'proceedingStatus':entry.get("proceedingStatus", ""),
                                'isSubmitted':entry.get("isSubmitted", ""),
                                'respStatus':entry.get("respStatus", ""),
                                'respId':entry.get("respId", ""),
                                'commType':entry.get("commType", ""),
                                'readFlag':entry.get("readFlag", ""),
                                'isRevisedItr':entry.get("isRevisedItr", ""),
                                'procdngModName':entry.get("procdngModName", "")or "Null",
                                'vcEnableFlag':entry.get("vcEnableFlag", "")or "Null",
                                'isActiveAR':entry.get("isActiveAR"),
                                'returnEverified':entry.get("returnEverified"), 
                                'discardAllowed':entry.get("discardAllowed", ""),
                                'documentCode':entry.get("documentCode", ""),
                                'isFileAppeal':entry.get("isFileAppeal", ""),
                                'isRectification':entry.get("isRectification", ""),
                                }
                                data_list.append(fya_all_notices)

                            Information = pd.DataFrame(data_list)
                            Information.replace('\x00', '', regex=True, inplace=True)

                            csv_file_path = fya_all_notices_path
                            file_exists = os.path.exists(csv_file_path)

                            curr_ef_db.execute("""
                                SELECT EXISTS (
                                    SELECT 1
                                    FROM information_schema.tables
                                    WHERE table_name = 'fya_all_notices'
                                );
                            """)

                            table_exists = curr_ef_db.fetchone()[0]

                            if not table_exists:
                                print("Table does not exist. Creating table...")
                                curr_ef_db.execute("""
                                    CREATE TABLE fya_all_notices (
                                        proceedingReqId text,
                                        pan text,
                                        nameOfAssesse text,
                                        headerSeqNo text,
                                        proceedingName text,
                                        financialYr text,
                                        proceedingLimitationDate text,
                                        proceedingLimitationDateTime text,
                                        proceedingType text,
                                        documentIdentificationNumber text,
                                        ay text,
                                        noticeSection text,
                                        description text,
                                        issuedOn text,
                                        issudedOnTime text,
                                        servedOn text,
                                        servedOnTime text,
                                        responseDueDate text,
                                        responseDueDateTime text,
                                        lastResponseSubmittedOn text,
                                        lastResponseSubmittedOnTime text,
                                        responseViewedByAoOn text,
                                        documentReferenceId text,
                                        proceedingStatus text,
                                        isSubmitted text,
                                        respStatus text,
                                        respId text,
                                        commType text,
                                        readFlag text,
                                        isRevisedItr text,
                                        procdngModName text,
                                        vcEnableFlag text,
                                        isActiveAR text,
                                        returnEverified text, 
                                        discardAllowed text,
                                        documentCode text,
                                        isFileAppeal text,
                                        isRectification text
                                    );
                                """)

                                conn_ef_db.commit()
                                print("Table 'fya_all_notices' created successfully.")
                                print()

                            if file_exists:
                                print("File already exists.")

                                if file_exists:
                                    existing_data = pd.read_csv(csv_file_path)
                                    existing_users = existing_data['proceedingReqId'].tolist()

                                    for index, row in Information.iterrows():
                                        if row['proceedingReqId'] in existing_users:
                                            print(f"FYA ALL NOTICES: {row['pan']} {row['proceedingReqId']} ALREADY EXISTED.")

                                        else:
                                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                                print(f"FYA ALL NOTICE: {row['pan']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")

                                            curr_ef_db.execute("""
                                            INSERT INTO fya_all_notices (
                                                proceedingReqId, pan, nameOfAssesse, headerSeqNo, proceedingName, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, proceedingType, documentIdentificationNumber, ay, noticeSection, 
                                                description, issuedOn, issudedOnTime, servedOn, servedOnTime, responseDueDate, responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, documentReferenceId, 
                                                proceedingStatus, isSubmitted, respStatus, respId, commType, readFlag, isRevisedItr, procdngModName, vcEnableFlag, isActiveAR, returnEverified,  discardAllowed, documentCode, isFileAppeal, isRectification
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                            """, (
                                                row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['headerSeqNo'], row['proceedingName'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'], row['proceedingType'], row['documentIdentificationNumber'], row['ay'], row['noticeSection'],
                                                row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'], row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['documentReferenceId'],
                                                row['proceedingStatus'], row['isSubmitted'], row['respStatus'], row['respId'], row['commType'], row['readFlag'], row['isRevisedItr'], row['procdngModName'], row['vcEnableFlag'], row['isActiveAR'], row['returnEverified'], row['discardAllowed'], row['documentCode'], row['isFileAppeal'], row['isRectification']
                                            ))

                                            conn_ef_db.commit()
                                            print(f"FYA All Notices: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                            print()

                            else:
                                Information.to_csv(csv_file_path, header=True, index=False)
                                print("FYA ALL NOTICE: CSV CREATED SUCCESSFULLY.")

                                for index, row in Information.iterrows():
                                    curr_ef_db.execute("""
                                    INSERT INTO fya_all_notices (
                                        proceedingReqId, pan, nameOfAssesse, headerSeqNo, proceedingName, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, proceedingType, documentIdentificationNumber, ay, noticeSection, 
                                        description, issuedOn, issudedOnTime, servedOn, servedOnTime, responseDueDate, responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, documentReferenceId, 
                                        proceedingStatus, isSubmitted, respStatus, respId, commType, readFlag, isRevisedItr, procdngModName, vcEnableFlag, isActiveAR, returnEverified,  discardAllowed, documentCode, isFileAppeal, isRectification
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['headerSeqNo'], row['proceedingName'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'], row['proceedingType'], row['documentIdentificationNumber'], row['ay'], row['noticeSection'],
                                        row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'], row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['documentReferenceId'],
                                        row['proceedingStatus'], row['isSubmitted'], row['respStatus'], row['respId'], row['commType'], row['readFlag'], row['isRevisedItr'], row['procdngModName'], row['vcEnableFlag'], row['isActiveAR'], row['returnEverified'], row['discardAllowed'], row['documentCode'], row['isFileAppeal'], row['isRectification']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"FYA All Notices: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()
                    
                except Exception as e:
                    print(f"{username} An error occurred for FYA all notices: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fya_all_notices
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYA ALL NOTICES: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA ALL NOTICES: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA ALL NOTICES: ERROR USERS CSV CREATED.")
                #------------------------------FYA NOTICE ALL NOTICES END------------------------------

                #------------------------------FYA NOTICE LETTER START------------------------------
                try:
                    time.sleep(3)
                    if response_fya.get('eproceedingRequests', []):

                        for proceeding_request in response_fya.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)

                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }
                            
                            response_view_notices = requests.request("POST", fya_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)
                            
                            view_notice = json.loads(response_view_notices.text)
                            
                            num_entries = len(view_notice)

                            data_list = []
                            for entry in view_notice:
                                
                                headerSeqNo = entry.get("headerSeqNo", "")
                                
                                payload_noticeletter = json.dumps({
                                "serviceName":"noticeletterpdf",
                                "headerSeqNo":headerSeqNo,
                                "procdngReqId":proceedingReqId,
                                "loggedInUserId":response_entity,
                                "header":{
                                    "formName":"FO-041_PCDNG"
                                    }
                                })
            
                                headers_noticeletter = {
                                'Content-Type': 'application/json',
                                'Cookie': 'AuthToken='+auth_token,
                                }
            
                                NoticeLetter = requests.request("POST", fya_notice_letter_url, headers = headers_noticeletter, data = payload_noticeletter)
                                time.sleep(3)

                                noticeletter = json.loads(NoticeLetter.text)

                                date_encrypt = noticeletter.get("date")
                                if date_encrypt is not None:
                                    dateDate = datetime.fromtimestamp(date_encrypt / 1000)
                                    date = dateDate.strftime('%d-%m-%Y')
                                    dateTime = dateDate.time()
                                else:
                                    date = "Null"
                                    dateTime = "Null"
                                
                                fya_notice_letter_list = {
                                'proceedingReqId': entry.get("proceedingReqId", "") or "Null",
                                "panNum":noticeletter.get("panNum", "") or "Null",
                                "userName":noticeletter.get("userName", ""),
                                "loggedInUserId":noticeletter.get("loggedInUserId", "") or "Null",
                                "noticeSection":noticeletter.get("noticeSection", "") or "Null",
                                "documentRefId":noticeletter.get("documentRefId", "") or "Null",
                                "description":noticeletter.get("description", "") or "Null",
                                "responseViewedByAO":noticeletter.get("responseViewedByAO", "") or "Null",
                                "proceedingName":noticeletter.get("proceedingName", "") or "Null",
                                "assessmentYear":noticeletter.get("assessmentYear", "") or "Null",
                                "noticeId":noticeletter.get("noticeId", "") or "Null",
                                "cc":noticeletter.get("cc", "") or "Null",
                                "mailBody":noticeletter.get("mailBody", "") or "Null",
                                "docNam":noticeletter.get("docNam", "") or "Null",
                                "headerSeqNo":noticeletter.get("headerSeqNo", "") or "Null",
                                "procdngReqId":noticeletter.get("procdngReqId", "") or "Null",
                                "applnId":noticeletter.get("applnId", "") or "Null",
                                "date":date,
                                "dateTime":dateTime,
                                "from":noticeletter.get("from", "") or "Null",
                                "subject":noticeletter.get("subject", "") or "Null",
                                "to":noticeletter.get("to", "") or "Null"
                                }
                                data_list.append(fya_notice_letter_list)

                            Information = pd.DataFrame(data_list)
                            Information.replace('\x00', '', regex=True, inplace=True)

                            csv_file_path = fya_notices_letter
                            file_exists = os.path.exists(csv_file_path)

                            curr_ef_db.execute("""
                                SELECT EXISTS (
                                    SELECT 1
                                    FROM information_schema.tables
                                    WHERE table_name = 'fya_notices_letter'
                                );
                            """)

                            table_exists = curr_ef_db.fetchone()[0]

                            if not table_exists:
                                print("Table does not exist. Creating table...")
                                curr_ef_db.execute("""
                                    CREATE TABLE fya_notices_letter (
                                        proceedingReqId text,
                                        panNum text,
                                        userName text,
                                        loggedInUserId text,
                                        noticeSection text,
                                        documentRefId text,
                                        description text,
                                        responseViewedByAO text,
                                        proceedingName text,
                                        assessmentYear text,
                                        noticeId text,
                                        cc text,
                                        mailBody text,
                                        docNam text,
                                        headerSeqNo text,
                                        procdngReqId text,
                                        applnId text,
                                        date text,
                                        dateTime text,
                                        frommail text,
                                        subject text,
                                        tomail text
                                    );
                                """)

                                conn_ef_db.commit()
                                print("Table 'fya_notices_letter' created successfully.")
                                print()

                            if file_exists:
                                print("File already exists.")

                                if file_exists:
                                    existing_data = pd.read_csv(csv_file_path)
                                    existing_users = existing_data['proceedingReqId'].tolist()

                                    for index, row in Information.iterrows():
                                        if row['proceedingReqId'] in existing_users:
                                            print(f"FYA NOTICE LETTER: {row['panNum']} {row['proceedingReqId']} ALREADY EXISTED.")

                                        else:
                                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                                print(f"FYA NOTICE LETTER: {row['panNum']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")

                                            curr_ef_db.execute("""
                                            INSERT INTO fya_notices_letter (
                                                proceedingReqId, panNum, userName, loggedInUserId, noticeSection, documentRefId, description, responseViewedByAO, proceedingName, 
                                                assessmentYear, noticeId, cc, mailBody, docNam, headerSeqNo, procdngReqId, applnId, date, dateTime, frommail, subject, tomail
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                            """, (
                                                row['proceedingReqId'], row['panNum'], row['userName'], row['loggedInUserId'], row['noticeSection'], row['documentRefId'], row['description'], row['responseViewedByAO'], row['proceedingName'],
                                                row['assessmentYear'], row['noticeId'], row['cc'], row['mailBody'], row['docNam'], row['headerSeqNo'], row['procdngReqId'], row['applnId'], row['date'], row['dateTime'], row['from'], row['subject'], row['to']
                                            ))

                                            conn_ef_db.commit()
                                            print(f"FYA Notice Letter: {row['panNum']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                            print()

                            else:
                                Information.to_csv(csv_file_path, header=True, index=False)
                                print("FYA NOTICE LETTER: CSV CREATED SUCCESSFULLY.")

                                for index, row in Information.iterrows():
                                    curr_ef_db.execute("""
                                    INSERT INTO fya_notices_letter (
                                        proceedingReqId, panNum, userName, loggedInUserId, noticeSection, documentRefId, description, responseViewedByAO, proceedingName, 
                                        assessmentYear, noticeId, cc, mailBody, docNam, headerSeqNo, procdngReqId, applnId, date, dateTime, frommail, subject, tomail
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['proceedingReqId'], row['panNum'], row['userName'], row['loggedInUserId'], row['noticeSection'], row['documentRefId'], row['description'], row['responseViewedByAO'], row['proceedingName'],
                                        row['assessmentYear'], row['noticeId'], row['cc'], row['mailBody'], row['docNam'], row['headerSeqNo'], row['procdngReqId'], row['applnId'], row['date'], row['dateTime'], row['from'], row['subject'], row['to']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"FYA Notice Letter: {row['panNum']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                except Exception as e:
                    print(f"{username} An error occurred for FYA notices letter: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fya_notice_letter
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYA NOTICE LETTER: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA NOTICE LETTER: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA NOTICE LETTER: ERROR USERS CSV CREATED.")
                #------------------------------FYA NOTICE LETTER END------------------------------

                #------------------------------FYA NOTICE DOWNLOAD START------------------------------
                try:
                    time.sleep(3)
                    if response_fya.get('eproceedingRequests', []):
                        for proceeding_request in response_fya.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)

                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }
                            
                            response_view_notices = requests.request("POST", fya_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)
                            
                            view_notice = json.loads(response_view_notices.text)

                            num_entries = len(view_notice)
                            
                            time.sleep(3)

                            for entry in view_notice:
                                headerSeqNo = entry.get("headerSeqNo", "")
                                payload_noticeDownload = json.dumps({
                                "serviceName": "noticeletterpdf",
                                "headerSeqNo": headerSeqNo,
                                "procdngReqId": proceedingReqId,
                                "loggedInUserId": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
            
                                headers_noticeDownload = {
                                'Content-Type': 'application/json',
                                'Cookie': 'AuthToken='+auth_token,
                                }
            
                                DownloadNotice = requests.request("POST", fya_download_notice_url, headers = headers_noticeDownload, data = payload_noticeDownload)
                                time.sleep(3)

                                downloadnotice = json.loads(DownloadNotice.text)
                                
                                satDocId = downloadnotice.get("satDocId")
                                doc_name = downloadnotice.get("docNam")

                                if satDocId:
                                    pdf_url = f"https://eportal.incometax.gov.in/iec/document/{satDocId}"

                                    headers_download = {
                                    'Cookie': 'AuthToken='+auth_token,
                                    'Referer': 'https://eportal.incometax.gov.in/iec/foservices/'
                                    }

                                    response = requests.get(pdf_url, headers=headers_download)

                                    if response.status_code == 200:
                                        file_path = os.path.join(fya_download_dir, f"{response_entity}_{doc_name}")
                                        
                                        with open(file_path, 'wb') as file:
                                            file.write(response.content)
                                        print(f'Document download for {response_entity}:', headerSeqNo)

                except Exception as e:
                    print(f"{username} An error occurred for FYA notice download: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fya_notice_download
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYA NOTICE DOWNLOAD: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA NOTICE DOWNLOAD: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA NOTICE DOWNLOAD: ERROR USERS CSV CREATED.")
                #------------------------------FYA NOTICE DOWNLOAD END------------------------------

        except Exception as e:
            print(f"{username} An error occurred for FYA: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_fya
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"FYA: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"FYA: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"FYA: ERROR USERS CSV CREATED.")
        #----------------------------------------FYA END------------------------------------------

        #----------------------------------------FYI START------------------------------------------
        try:
            time.sleep(3)
            payload_userInfo_1 = json.dumps({
            "serviceName": "userProfileService"
            })

            headers_userInfo_1 = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn': 'userProfileService'
            }

            payload_eproceedingCount = json.dumps({
            "serviceName": "getEProceedingsDetail",
            "pan": response_entity,
            "prcdngStatusFlag": "FYI",
            "header": {
                "formName": "FO-041_PCDNG"
            }
            })

            headers_eproceedingCount = {
            'Content-Type': 'application/json',
            'Cookie': 'AuthToken='+auth_token,
            'sn':'getEProceedingsDetail'
            }

            UserInfo_1 = requests.request("POST", userinfo_url, headers = headers_userInfo_1, data = payload_userInfo_1)
            Response_FYI = requests.request("POST", fyi_notice_count_url, headers = headers_eproceedingCount, data = payload_eproceedingCount)
            time.sleep(3)
        
            basic_info = json.loads(UserInfo_1.text)
            response_fyi = json.loads(Response_FYI.text)
        
            if UserInfo_1.status_code == 200 and Response_FYI.status_code == 200:
                #------------------------------FYI NOTICE COUNT START------------------------------
                try:
                    time.sleep(3)
                    fyi_count_list = [{
                    "userId": basic_info.get("userId", ""),
                    "roleDesc": basic_info.get("roleDesc", ""),
                    "orgName":basic_info.get("orgName", "")or "Null",
                    "contactFirstName":basic_info.get("contactFirstName", "")or "Null",
                    "contactMiddleName":basic_info.get("contactMiddleName", "")or "Null",
                    "contactLastName":basic_info.get("contactLastName", "")or "Null",
                    "firstName": basic_info.get("firstName", ""),
                    "midName": basic_info.get("midName", ""),
                    "lastName": basic_info.get("lastName", ""),
                    "eProceedingsInfoCount" : response_fyi.get("eProceedingsInfoCount", ""),
                    "eProceedingsForActionCount" : response_fyi.get("eProceedingsForActionCount", "")
                    }]
            
                    Information = pd.DataFrame(fyi_count_list)
                    Information.replace('\x00', '', regex=True, inplace=True)
                    
                    csv_file_path = fyi_notice_count
                    file_exists = os.path.exists(csv_file_path)

                    curr_ef_db.execute("""
                        SELECT EXISTS (
                            SELECT 1
                            FROM information_schema.tables
                            WHERE table_name = 'fyi_notice_count'
                        );
                    """)

                    table_exists = curr_ef_db.fetchone()[0]

                    if not table_exists:
                        print("Table does not exist. Creating table...")
                        curr_ef_db.execute("""
                            CREATE TABLE fyi_notice_count (
                                userId text,
                                roleDesc text,
                                orgName text,
                                contactFirstName text,
                                contactMiddleName text,
                                contactLastName text,
                                firstName text,
                                midName text,
                                lastName text,
                                eProceedingsInfoCount text,
                                eProceedingsForActionCount text
                            );
                        """)

                        conn_ef_db.commit()
                        print("Table 'fyi_notice_count' created successfully.")
                        print()

                    if file_exists:
                        print("File already exists.")

                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['userId'].tolist()

                            for index, row in Information.iterrows():
                                if row['userId'] in existing_users:
                                    print(f"FYI NOTICE COUNT: {row['userId']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYI NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY.")
                                    
                                    curr_ef_db.execute("""
                                    INSERT INTO fyi_notice_count (
                                        userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                        firstName, midName, lastName, eProceedingsInfoCount, eProceedingsForActionCount
                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                    """, (
                                        row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'],
                                        row['contactMiddleName'], row['contactLastName'], row['firstName'], row['midName'], 
                                        row['lastName'], row['eProceedingsInfoCount'], row['eProceedingsForActionCount']
                                    ))

                                    conn_ef_db.commit()
                                    print(f"FYI NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                    print()

                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print("FYI NOTICE COUNT: CSV CREATED SUCCESSFULLY.")

                        for index, row in Information.iterrows():
                            curr_ef_db.execute("""
                            INSERT INTO fyi_notice_count (
                                userId, roleDesc, orgName, contactFirstName, contactMiddleName, contactLastName, 
                                firstName, midName, lastName, eProceedingsInfoCount, eProceedingsForActionCount
                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['userId'], row['roleDesc'], row['orgName'], row['contactFirstName'],
                                row['contactMiddleName'], row['contactLastName'], row['firstName'], row['midName'], 
                                row['lastName'], row['eProceedingsInfoCount'], row['eProceedingsForActionCount']
                            ))

                            conn_ef_db.commit()
                            print(f"FYI NOTICE COUNT: {row['userId']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()

                except Exception as e:
                    print(f"{username} An error occurred for FYI Notice Count: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fyi_count
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI NOTICE COUNT: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYI NOTICE COUNT: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYI NOTICE COUNT: ERROR USERS CSV CREATED.")
                #------------------------------FYI NOTICE COUNT END------------------------------

                #------------------------------FYI NOTICE DESCRIPTION START------------------------------
                try:
                    time.sleep(3)
                    if response_fyi.get('eproceedingRequests', []):
                        eprorequest = response_fyi.get('eproceedingRequests', [])

                        data_list = []
                        for entry in eprorequest:

                            proceedingLimitationDate_encrypt = entry.get("proceedingLimitationDate", "")
                            if proceedingLimitationDate_encrypt:
                                proceedingLimitationDate_Date = datetime.fromtimestamp(proceedingLimitationDate_encrypt / 1000)
                                proceedingLimitationDate = proceedingLimitationDate_Date.strftime('%d-%m-%Y')
                                proceedingLimitationDateTime = proceedingLimitationDate_Date.time()
                            else:
                                proceedingLimitationDate = "Null"
                                proceedingLimitationDateTime = "Null"

                            responseDate_encrypt = entry.get("responseDate", "")
                            if responseDate_encrypt:
                                responseDate_Date = datetime.fromtimestamp(responseDate_encrypt / 1000)
                                responseDate = responseDate_Date.strftime('%d-%m-%Y')
                                responseDateTime = responseDate_Date.time()
                            else:
                                responseDate = "Null"
                                responseDateTime = "Null"

                            acknowledgementNo_encrypt =  entry.get("acknowledgementNo", "")
                            if acknowledgementNo_encrypt:
                                acknowledgementNo = f"'{acknowledgementNo_encrypt}"
                            else:
                                acknowledgementNo = "Null"


                            issuedOn_encrypt = entry.get("issuedOn", "")
                            if issuedOn_encrypt:
                                issuedOn_Date = datetime.fromtimestamp(issuedOn_encrypt / 1000)
                                issuedOn = issuedOn_Date.strftime('%d-%m-%Y')
                                issudedOnTime = issuedOn_Date.time()
                            else:
                                issuedOn = "Null"
                                issuedOnTime = "Null"
                            
            
                            servedOn_encrypt = entry.get("servedOn", "")
                            if servedOn_encrypt:
                                servedOn_Date = datetime.fromtimestamp(servedOn_encrypt / 1000)
                                servedOn = servedOn_Date.strftime('%d-%m-%Y')
                                servedOnTime = servedOn_Date.time()
                            else:
                                servedOn = "Null"
                                servedOnTime = "Null"
            
                            responseDueDate_encrypt = entry.get("responseDueDate", "")
                            if responseDueDate_encrypt:
                                responseDueDate_Date = datetime.fromtimestamp(responseDueDate_encrypt / 1000)
                                responseDueDate = responseDueDate_Date.strftime('%d-%m-%Y')
                                responseDueDateTime = responseDueDate_Date.time()
                            else:
                                responseDueDate = "Null"
                                responseDueDateTime = "Null"
                            
                            lastResponseSubmittedOn_encrypt = entry.get("lastResponseSubmittedOn", "")
                            if lastResponseSubmittedOn_encrypt:
                                lastResponseSubmittedOn_Date = datetime.fromtimestamp(lastResponseSubmittedOn_encrypt / 1000)
                                lastResponseSubmittedOn = lastResponseSubmittedOn_Date.strftime('%d-%m-%Y')
                                lastResponseSubmittedOnTime = lastResponseSubmittedOn_Date.time()
                            else:
                                lastResponseSubmittedOn = "Null"
                                lastResponseSubmittedOnTime = "Null"

                            proceedingClosureDate_encrypt = entry.get("proceedingClosureDate", "")
                            if proceedingClosureDate_encrypt:
                                proceedingClosureDate_Date = datetime.fromtimestamp(proceedingClosureDate_encrypt / 1000)
                                proceedingClosureDate = proceedingClosureDate_Date.strftime('%d-%m-%Y')
                                proceedingClosureDateTime = proceedingClosureDate_Date.time()
                            else:
                                proceedingClosureDate = "Null"
                                proceedingClosureDateTime = "Null"

                        
                            fyi_notice_description_list = {
                            'proceedingReqId': entry.get("proceedingReqId", ""),
                            'pan': entry.get("pan", ""),
                            'nameOfAssesse': entry.get("nameOfAssesse", ""),
                            'proceedingName':entry.get("proceedingName", ""),
                            'itrType':entry.get("itrType", ""),
                            'assessmentYear':entry.get("assessmentYear", ""),
                            'financialYr':entry.get("financialYr", ""),
                            'proceedingLimitationDate':proceedingLimitationDate,
                            'proceedingLimitationDateTime':proceedingLimitationDateTime,
                            'noticeName':entry.get("noticeName", ""),
                            'responseDate':responseDate,
                            'responseDateTime':responseDateTime,
                            'acknowledgementNo':acknowledgementNo,
                            'viewNoticeCount':entry.get("viewNoticeCount", ""),
                            'proceedingType':entry.get("proceedingType", ""),
                            'issuedOn':issuedOn,
                            'issuedOnTime':issuedOnTime,
                            "servedOn" : servedOn,
                            "servedOnTime":servedOnTime,
                            "responseDueDate" : responseDueDate,
                            "responseDueDateTime":responseDueDateTime,
                            "lastResponseSubmittedOn" : lastResponseSubmittedOn,
                            "lastResponseSubmittedOnTime":lastResponseSubmittedOnTime,
                            'responseViewedByAoOn':entry.get("responseViewedByAoOn", ""),
                            'proceedingClosureDate':proceedingClosureDate,
                            'proceedingClosureDateTime':proceedingClosureDateTime,
                            'proceedingClosureOrder':entry.get("proceedingClosureOrder", "") or "Null",
                            'proceedingStatus':entry.get("proceedingStatus", ""),
                            'respStatus':entry.get("respStatus", ""),
                            'respId':entry.get("respId", ""),
                            'commType':entry.get("commType", ""),
                            'readFlag':entry.get("readFlag", ""),
                            'facelessFlag':entry.get("facelessFlag", ""),
                            'returnEverified':entry.get("returnEverified", "")or "Null",
                            'discardAllowed':entry.get("discardAllowed", "")or "Null",
                            'new':entry.get("new", ""),
                            'isFileAppeal':entry.get("isFileAppeal", ""),
                            'isRectification':entry.get("isRectification", ""),
                            }
                            data_list.append(fyi_notice_description_list)

                        Information = pd.DataFrame(data_list)
                        Information.replace('\x00', '', regex=True, inplace=True)
                        
                        csv_file_path = fyi_notice_description
                        file_exists = os.path.exists(csv_file_path)

                        curr_ef_db.execute("""
                            SELECT EXISTS (
                                SELECT 1
                                FROM information_schema.tables
                                WHERE table_name = 'fyi_notice_description'
                            );
                        """)

                        table_exists = curr_ef_db.fetchone()[0]

                        if not table_exists:
                            print("Table does not exist. Creating table...")
                            curr_ef_db.execute("""
                                CREATE TABLE fyi_notice_description (
                                    proceedingReqId text,
                                    pan text,
                                    nameOfAssesse text,
                                    proceedingName text,
                                    itrType text,
                                    assessmentYear text,
                                    financialYr text,
                                    proceedingLimitationDate text,
                                    proceedingLimitationDateTime text,
                                    noticeName text,
                                    responseDate text,
                                    responseDateTime text,
                                    acknowledgementNo text,
                                    viewNoticeCount text,
                                    proceedingType text,
                                    issuedOn text,
                                    issuedOnTime text,
                                    servedOn text,
                                    servedOnTime text,
                                    responseDueDate text,
                                    responseDueDateTime text,
                                    lastResponseSubmittedOn text,
                                    lastResponseSubmittedOnTime text,
                                    responseViewedByAoOn text,
                                    proceedingClosureDate text,
                                    proceedingClosureDateTime text,
                                    proceedingClosureOrder text,
                                    proceedingStatus text,
                                    respStatus text,
                                    respId text,
                                    commType text,
                                    readFlag text,
                                    facelessFlag text,
                                    returnEverified text,
                                    discardAllowed text,
                                    new text,
                                    isFileAppeal text,
                                    isRectification text
                                );
                            """)

                            conn_ef_db.commit()
                            print("Table 'fyi_notice_description' created successfully.")
                            print()

                        if file_exists:
                            print("File already exists.")

                            if file_exists:
                                existing_data = pd.read_csv(csv_file_path)
                                existing_users = existing_data['proceedingReqId'].tolist()

                                for index, row in Information.iterrows():
                                    if row['proceedingReqId'] in existing_users:
                                        print(f"FYI NOTICE DESCRIPTION: {row['pan']} {row['proceedingReqId']} ALREADY EXISTED.")

                                    else:
                                        with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                            Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                            print(f"FYI NOTICE DESCRIPTION: {row['pan']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")

                                        
                                        curr_ef_db.execute("""
                                            INSERT INTO fyi_notice_description (
                                                proceedingReqId, pan, nameOfAssesse, proceedingName, itrType, assessmentYear, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, noticeName, 
                                                responseDate, responseDateTime, acknowledgementNo, viewNoticeCount, proceedingType, issuedOn, issuedOnTime, servedOn, servedOnTime, responseDueDate, responseDueDateTime, 
                                                lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, proceedingClosureDate, proceedingClosureDateTime, proceedingClosureOrder, proceedingStatus, 
                                                respStatus, respId, commType, readFlag, facelessFlag, returnEverified, discardAllowed, new, isFileAppeal, isRectification
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                        """, (
                                            row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['proceedingName'], row['itrType'], row['assessmentYear'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'],
                                            row['noticeName'], row['responseDate'], row['responseDateTime'], row['acknowledgementNo'],  row['viewNoticeCount'], row['proceedingType'], row['issuedOn'], row['issuedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                            row['responseDueDateTime'], row['lastResponseSubmittedOn'],row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['proceedingClosureDate'],row['proceedingClosureDateTime'], row['proceedingClosureOrder'],
                                            row['proceedingStatus'], row['respStatus'], row['respId'], row['commType'], row['readFlag'],row['facelessFlag'], row['returnEverified'], row['discardAllowed'], row['new'],row['isFileAppeal'], row['isRectification']
                                        ))

                                        conn_ef_db.commit()
                                        print(f"FYI Notice Description: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                        print()

                        else:
                            Information.to_csv(csv_file_path, header=True, index=False)
                            print("FYI NOTICE DESCRIPTION: CSV CREATED SUCCESSFULLY.")

                            for index, row in Information.iterrows():
                                curr_ef_db.execute("""
                                INSERT INTO fyi_notice_description (
                                    proceedingReqId, pan, nameOfAssesse, proceedingName, itrType, assessmentYear, financialYr, proceedingLimitationDate, proceedingLimitationDateTime, noticeName, 
                                    responseDate, responseDateTime, acknowledgementNo, viewNoticeCount, proceedingType, issuedOn, issuedOnTime, servedOn, servedOnTime, responseDueDate, responseDueDateTime, 
                                    lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, proceedingClosureDate, proceedingClosureDateTime, proceedingClosureOrder, proceedingStatus, 
                                    respStatus, respId, commType, readFlag, facelessFlag, returnEverified, discardAllowed, new, isFileAppeal, isRectification
                                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                            """, (
                                row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['proceedingName'], row['itrType'], row['assessmentYear'], row['financialYr'], row['proceedingLimitationDate'], row['proceedingLimitationDateTime'],
                                row['noticeName'], row['responseDate'], row['responseDateTime'], row['acknowledgementNo'],  row['viewNoticeCount'], row['proceedingType'], row['issuedOn'], row['issuedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                row['responseDueDateTime'], row['lastResponseSubmittedOn'],row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], row['proceedingClosureDate'],row['proceedingClosureDateTime'], row['proceedingClosureOrder'],
                                row['proceedingStatus'], row['respStatus'], row['respId'], row['commType'], row['readFlag'],row['facelessFlag'], row['returnEverified'], row['discardAllowed'], row['new'],row['isFileAppeal'], row['isRectification']
                            ))

                            conn_ef_db.commit()
                            print(f"FYI Notice Description: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                            print()
                
                except Exception as e:
                    print(f"{username} An error occurred for FYI Notice Description: {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fyi_notice_description
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI NOTICE DESCRIPTION: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA NOTICE DESCRIPTION: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA NOTICE DESCRIPTION: ERROR USERS CSV CREATED.")
                #------------------------------FYI NOTICE DESCRIPTION END------------------------------

                #------------------------------FYI ALL NOTICES START------------------------------
                try:
                    time.sleep(3)
                    if response_fyi.get('eproceedingRequests', []):
                        for proceeding_request in response_fyi.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)
                        
                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }

                            response_view_notices = requests.request("POST", fyi_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)

                            if response_view_notices.status_code==200:
                                try:
                                    view_notice = json.loads(response_view_notices.text)

                                    num_entries = len(view_notice)

                                    data_list = []
                                    for entry in view_notice:

                                        documentIdentificationNumber_encrypt =  entry.get("documentIdentificationNumber", "")
                                        if documentIdentificationNumber_encrypt is not None and documentIdentificationNumber_encrypt != "":
                                            if isinstance(documentIdentificationNumber_encrypt, str):
                                                documentIdentificationNumber = documentIdentificationNumber_encrypt
                                            else:
                                                documentIdentificationNumber = f"'{documentIdentificationNumber_encrypt}"
                                        else:
                                            documentIdentificationNumber = "Null"

                                        proceedingLimitationDate_encrypt = entry.get("proceedingLimitationDate", "")
                                        if proceedingLimitationDate_encrypt is not None and proceedingLimitationDate_encrypt in entry:
                                            proceedingLmtDate = datetime.fromtimestamp(proceedingLimitationDate_encrypt / 1000)
                                            proceedingLimitationDate = proceedingLmtDate.strftime('%d-%m-%Y')
                                            proceedingLimitationDateTime = proceedingLmtDate.time()
                                        else:
                                            proceedingLimitationDate = "Null"
                                            proceedingLimitationDateTime = "Null"
                                        
                                        issuedOn_encrypt = entry.get("issuedOn", "")
                                        if issuedOn_encrypt is not None:
                                            issuedDate = datetime.fromtimestamp(issuedOn_encrypt / 1000)
                                            issuedOn = issuedDate.strftime('%d-%m-%Y')
                                            issudedOnTime = issuedDate.time()
                                        else:
                                            issuedOn = "Null"
                                            issudedOnTime = "Null"

                                        servedOn_encrypt = entry.get("servedOn", "")
                                        if servedOn_encrypt is not None:
                                            servedDate = datetime.fromtimestamp(servedOn_encrypt / 1000)
                                            servedOn = servedDate.strftime('%d-%m-%Y')
                                            servedOnTime = servedDate.time()
                                        else:
                                            servedOn = "Null"
                                            servedOnTime = "Null"

                                        responseDueDate_encrypt = entry.get("responseDueDate", "")
                                        if responseDueDate_encrypt is not None:
                                            responseDate = datetime.fromtimestamp(responseDueDate_encrypt / 1000)
                                            responseDueDate = responseDate.strftime('%d-%m-%Y')
                                            responseDueDateTime = responseDate.time()
                                        else:
                                            responseDueDate = "Null"
                                            responseDueDateTime = "Null"

                                        lastResponseSubmittedOn_encrypt = entry.get("lastResponseSubmittedOn", "")
                                        if lastResponseSubmittedOn_encrypt is not None:
                                            lastResponseDate = datetime.fromtimestamp(lastResponseSubmittedOn_encrypt / 1000)
                                            lastResponseSubmittedOn = lastResponseDate.strftime('%d-%m-%Y')
                                            lastResponseSubmittedOnTime = lastResponseDate.time()
                                        else:
                                            lastResponseSubmittedOn = "Null"
                                            lastResponseSubmittedOnTime = "Null"
                                        
                                        user_basic_info_list = {
                                            'proceedingReqId': entry.get("proceedingReqId", ""),
                                            'pan': entry.get("pan", ""),
                                            'nameOfAssesse': entry.get("nameOfAssesse", ""),
                                            'headerSeqNo': entry.get("headerSeqNo", ""),
                                            'proceedingName':entry.get("proceedingName", ""),
                                            'financialYr':entry.get("financialYr", ""),
                                            'proceedingLimitationDate':proceedingLimitationDate,
                                            'proceedingLimitationDateTime':proceedingLimitationDateTime,
                                            'proceedingType':entry.get("proceedingType", ""),
                                            'documentIdentificationNumber':documentIdentificationNumber,
                                            'ay':entry.get("ay", ""),
                                            'noticeSection':entry.get("noticeSection", ""),
                                            'description':entry.get("description", ""),
                                            "issuedOn" : issuedOn,
                                            "issudedOnTime": issudedOnTime,
                                            "servedOn" : servedOn,
                                            "servedOnTime":servedOnTime,
                                            "responseDueDate" : responseDueDate,
                                            "responseDueDateTime":responseDueDateTime,
                                            "lastResponseSubmittedOn" : lastResponseSubmittedOn,
                                            "lastResponseSubmittedOnTime":lastResponseSubmittedOnTime,
                                            'responseViewedByAoOn':entry.get("responseViewedByAoOn", ""),
                                            'documentReferenceId':entry.get("documentReferenceId", ""),
                                            'proceedingStatus':entry.get("proceedingStatus", ""),
                                            'isSubmitted':entry.get("isSubmitted", ""),
                                            'respStatus':entry.get("respStatus", ""),
                                            'respId':entry.get("respId", ""),
                                            'commType':entry.get("commType", ""),
                                            'readFlag':entry.get("readFlag", ""),
                                            'isRevisedItr':entry.get("isRevisedItr", ""),
                                            'procdngModName':entry.get("procdngModName", "")or "Null",
                                            'vcEnableFlag':entry.get("vcEnableFlag", "")or "Null",
                                            'returnEverified':entry.get("returnEverified", ""),
                                            'discardAllowed':entry.get("discardAllowed", ""),
                                            'documentCode':entry.get("documentCode", ""),
                                            'isFileAppeal':entry.get("isFileAppeal", ""),
                                            'isRectification':entry.get("isRectification", ""),
                                        }
                                        data_list.append(user_basic_info_list)

                                    Information = pd.DataFrame(data_list)
                                    Information.replace('\x00', '', regex=True, inplace=True)

                                    csv_file_path = fyi_all_notices
                                    file_exists = os.path.exists(csv_file_path)

                                    curr_ef_db.execute("""
                                        SELECT EXISTS (
                                            SELECT 1
                                            FROM information_schema.tables
                                            WHERE table_name = 'fyi_all_notices'
                                        );
                                    """)

                                    table_exists = curr_ef_db.fetchone()[0]

                                    if not table_exists:
                                        print("Table does not exist. Creating table...")
                                        curr_ef_db.execute("""
                                            CREATE TABLE fyi_all_notices (
                                                proceedingReqId text,
                                                pan text,
                                                nameOfAssesse text,
                                                headerSeqNo text,
                                                proceedingName text,
                                                financialYr text,
                                                proceedingLimitationDate text,
                                                proceedingLimitationDateTime text,
                                                proceedingType text,
                                                documentIdentificationNumber text,
                                                ay text,
                                                noticeSection text,
                                                description text,
                                                issuedOn text,
                                                issudedOnTime text,
                                                servedOn text,
                                                servedOnTime text,
                                                responseDueDate text,
                                                responseDueDateTime text,
                                                lastResponseSubmittedOn text,
                                                lastResponseSubmittedOnTime text,
                                                responseViewedByAoOn text,
                                                documentReferenceId text,
                                                proceedingStatus text,
                                                isSubmitted text,
                                                respStatus text,
                                                respId text,
                                                commType text,
                                                readFlag text,
                                                isRevisedItr text,
                                                procdngModName text,
                                                vcEnableFlag text,
                                                returnEverified text,
                                                discardAllowed text,
                                                documentCode text,
                                                isFileAppeal text,
                                                isRectification text
                                            );
                                        """)

                                        conn_ef_db.commit()
                                        print("Table 'fyi_all_notices' created successfully.")
                                        print()

                                    if file_exists:
                                        print("File already exists.")

                                        if file_exists:
                                            existing_data = pd.read_csv(csv_file_path)
                                            existing_users = existing_data['proceedingReqId'].tolist()

                                            for index, row in Information.iterrows():
                                                if row['proceedingReqId'] in existing_users:
                                                    print(f"FYI ALL NOTICES: {row['pan']} {row['proceedingReqId']} ALREADY EXISTED.")

                                                else:
                                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                                        print(f"FYI ALL NOTICE: {row['pan']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")

                                                    curr_ef_db.execute("""
                                                    INSERT INTO fyi_all_notices (
                                                        proceedingReqId, pan, nameOfAssesse, headerSeqNo, proceedingName, financialYr, 
                                                        proceedingLimitationDate, proceedingLimitationDateTime, proceedingType, documentIdentificationNumber, 
                                                        ay, noticeSection, description, issuedOn, issudedOnTime, servedOn, servedOnTime, 
                                                        responseDueDate, responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, 
                                                        documentReferenceId, proceedingStatus, isSubmitted, respStatus, respId, commType, readFlag, 
                                                        isRevisedItr, procdngModName, vcEnableFlag, returnEverified, discardAllowed, documentCode, 
                                                        isFileAppeal, isRectification
                                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                                    """, (
                                                        row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['headerSeqNo'], row['proceedingName'], row['financialYr'], 
                                                        row['proceedingLimitationDate'], row['proceedingLimitationDateTime'], row['proceedingType'], row['documentIdentificationNumber'], 
                                                        row['ay'], row['noticeSection'], row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], 
                                                        row['responseDueDate'], row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], 
                                                        row['documentReferenceId'], row['proceedingStatus'], row['isSubmitted'], row['respStatus'], row['respId'], row['commType'], row['readFlag'], 
                                                        row['isRevisedItr'], row['procdngModName'], row['vcEnableFlag'], row['returnEverified'], row['discardAllowed'], row['documentCode'], 
                                                        row['isFileAppeal'], row['isRectification']
                                                    ))

                                                    conn_ef_db.commit()
                                                    print(f"FYI All Notices: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                                    print()

                                    else:
                                        Information.to_csv(csv_file_path, header=True, index=False)
                                        print("FYI ALL NOTICE: CSV CREATED SUCCESSFULLY.")

                                        for index, row in Information.iterrows():
                                            curr_ef_db.execute("""
                                            INSERT INTO fyi_all_notices (
                                                proceedingReqId, pan, nameOfAssesse, headerSeqNo, proceedingName, financialYr, 
                                                proceedingLimitationDate, proceedingLimitationDateTime, proceedingType, documentIdentificationNumber, 
                                                ay, noticeSection, description, issuedOn, issudedOnTime, servedOn, servedOnTime, 
                                                responseDueDate, responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, responseViewedByAoOn, 
                                                documentReferenceId, proceedingStatus, isSubmitted, respStatus, respId, commType, readFlag, 
                                                isRevisedItr, procdngModName, vcEnableFlag, returnEverified, discardAllowed, documentCode, 
                                                isFileAppeal, isRectification
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                            """, (
                                                row['proceedingReqId'], row['pan'], row['nameOfAssesse'], row['headerSeqNo'], row['proceedingName'], row['financialYr'], 
                                                row['proceedingLimitationDate'], row['proceedingLimitationDateTime'], row['proceedingType'], row['documentIdentificationNumber'], 
                                                row['ay'], row['noticeSection'], row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], 
                                                row['responseDueDate'], row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'], row['responseViewedByAoOn'], 
                                                row['documentReferenceId'], row['proceedingStatus'], row['isSubmitted'], row['respStatus'], row['respId'], row['commType'], row['readFlag'], 
                                                row['isRevisedItr'], row['procdngModName'], row['vcEnableFlag'], row['returnEverified'], row['discardAllowed'], row['documentCode'], 
                                                row['isFileAppeal'], row['isRectification']
                                            ))

                                            conn_ef_db.commit()
                                            print(f"FYI All Notices: {row['pan']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                            print()

                                except Exception as e:
                                    print(f"No Other Notice for {response_entity}")

                except Exception as e:
                    print(f"{username} An error occurred for FYI All Notices : {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fyi_all_notices
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI ALL NOTICES: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYA ALL NOTICES: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYA ALL NOTICE: ERROR USERS CSV CREATED.")
                #------------------------------FYI ALL NOTICES END------------------------------

                #------------------------------FYI NOTICES LETTER START------------------------------
                try:
                    time.sleep(3)
                    if response_fyi.get('eproceedingRequests', []):
                        for proceeding_request in response_fyi.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)

                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }
                            
                            response_view_notices = requests.request("POST", fyi_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)

                            if response_view_notices.status_code==200:
                                try:
                                    time.sleep(3)
                                    view_notice = json.loads(response_view_notices.text)
                                    
                                    num_entries = len(view_notice)

                                    data_list = []
                                    for entry in view_notice:
                                        
                                        headerSeqNo = entry.get("headerSeqNo", "")
                                        
                                        payload_noticeletter = json.dumps({
                                        "serviceName":"noticeletterpdf",
                                        "headerSeqNo":headerSeqNo,
                                        "procdngReqId":proceedingReqId,
                                        "loggedInUserId":response_entity,
                                        "header":{
                                            "formName":"FO-041_PCDNG"
                                            }
                                        })
                    
                                        headers_noticeletter = {
                                        'Content-Type': 'application/json',
                                        'Cookie': 'AuthToken='+auth_token,
                                        }

                                        NoticeLetter = requests.request("POST", fyi_notice_letter_url, headers = headers_noticeletter, data = payload_noticeletter)
                                        time.sleep(3)

                                        noticeletter = json.loads(NoticeLetter.text)

                                        issuedOn_encrypt = noticeletter.get("issuedOn", "")
                                        if issuedOn_encrypt is not None:
                                            issuedDate = datetime.fromtimestamp(issuedOn_encrypt / 1000)
                                            issuedOn = issuedDate.strftime('%d-%m-%Y')
                                            issudedOnTime = issuedDate.time()
                                        else:
                                            issuedOn = "Null"
                                            issudedOnTime = "Null"

                                        servedOn_encrypt = noticeletter.get("servedOn", "")
                                        if servedOn_encrypt is not None:
                                            servedDate = datetime.fromtimestamp(servedOn_encrypt / 1000)
                                            servedOn = servedDate.strftime('%d-%m-%Y')
                                            servedOnTime = servedDate.time()
                                        else:
                                            servedOn = "Null"
                                            servedOnTime = "Null"

                                        responseDueDate_encrypt = noticeletter.get("responseDueDate", "")
                                        if responseDueDate_encrypt is not None:
                                            responseDate = datetime.fromtimestamp(responseDueDate_encrypt / 1000)
                                            responseDueDate = responseDate.strftime('%d-%m-%Y')
                                            responseDueDateTime = responseDate.time()
                                        else:
                                            responseDueDate = "Null"
                                            responseDueDateTime = "Null"

                                        lastResponseSubmittedOn_encrypt = noticeletter.get("lastResponseSubmittedOn", "")
                                        if lastResponseSubmittedOn_encrypt is not None:
                                            lastResponseDate = datetime.fromtimestamp(lastResponseSubmittedOn_encrypt / 1000)
                                            lastResponseSubmittedOn = lastResponseDate.strftime('%d-%m-%Y')
                                            lastResponseSubmittedOnTime = lastResponseDate.time()
                                        else:
                                            lastResponseSubmittedOn = "Null"
                                            lastResponseSubmittedOnTime = "Null"

                                        date_encrypt = noticeletter.get("date", "")
                                        if date_encrypt is not None:
                                            dateDate = datetime.fromtimestamp(date_encrypt / 1000)
                                            date = dateDate.strftime('%d-%m-%Y')
                                            dateTime = dateDate.time()
                                        else:
                                            date = "Null"
                                            dateTime = "Null"
                                        
                                        user_basic_info_list = {
                                        'proceedingReqId': entry.get("proceedingReqId", ""),
                                        "panNum":noticeletter.get("panNum", ""),
                                        "userName":noticeletter.get("userName", ""),
                                        "loggedInUserId":noticeletter.get("loggedInUserId", "")or "Null",
                                        "noticeSection":noticeletter.get("noticeSection", ""),
                                        "documentRefId":noticeletter.get("documentRefId", "")or "Null",
                                        "description":noticeletter.get("description", "")or "Null",
                                        "issuedOn":issuedOn,
                                        "issudedOnTime":issudedOnTime,
                                        "servedOn":servedOn,
                                        "servedOnTime":servedOnTime,    
                                        "responseDueDate":responseDueDate,
                                        "responseDueDateTime":responseDueDateTime,
                                        "lastResponseSubmittedOn":lastResponseSubmittedOn,
                                        "lastResponseSubmittedOnTime":lastResponseSubmittedOnTime,
                                        "responseViewedByAO":noticeletter.get("responseViewedByAO", "") or "Null",
                                        "proceedingName":noticeletter.get("proceedingName", ""),
                                        "assessmentYear":noticeletter.get("assessmentYear", ""),
                                        "noticeId":noticeletter.get("noticeId", ""),
                                        "cc":noticeletter.get("cc", ""),
                                        "mailBody":noticeletter.get("mailBody", ""),
                                        "docNam":noticeletter.get("docNam", ""),
                                        "headerSeqNo":noticeletter.get("headerSeqNo", ""),
                                        "procdngReqId":noticeletter.get("procdngReqId", "") or "Null",
                                        "applnId":noticeletter.get("applnId", "")or "Null",
                                        "date":date,
                                        "dateTime":dateTime,
                                        "from":noticeletter.get("from", ""),
                                        "subject":noticeletter.get("subject", ""),
                                        "to":noticeletter.get("to", ""),
                                        }
                                        data_list.append(user_basic_info_list)

                                    Information = pd.DataFrame(data_list)
                                    Information.replace('\x00', '', regex=True, inplace=True)
                                    
                                    csv_file_path = fyi_notices_letter
                                    file_exists = os.path.exists(csv_file_path)

                                    curr_ef_db.execute("""
                                        SELECT EXISTS (
                                            SELECT 1
                                            FROM information_schema.tables
                                            WHERE table_name = 'fyi_notices_letter'
                                        );
                                    """)

                                    table_exists = curr_ef_db.fetchone()[0]

                                    if not table_exists:
                                        print("Table does not exist. Creating table...")
                                        curr_ef_db.execute("""
                                            CREATE TABLE fyi_notices_letter (
                                                proceedingReqId text,
                                                panNum text,
                                                userName text,
                                                loggedInUserId text,
                                                noticeSection text,
                                                documentRefId text,
                                                description text,
                                                issuedOn text,
                                                issudedOnTime text,
                                                servedOn text,
                                                servedOnTime text,
                                                responseDueDate text,
                                                responseDueDateTime text,
                                                lastResponseSubmittedOn text,
                                                lastResponseSubmittedOnTime text,
                                                responseViewedByAO text,
                                                proceedingName text,
                                                assessmentYear text,
                                                noticeId text,
                                                cc text,
                                                mailBody text,
                                                docNam text,
                                                headerSeqNo text,
                                                procdngReqId text,
                                                applnId text,
                                                date text,
                                                dateTime text,
                                                frommail text,
                                                subject text,
                                                tomail text
                                            );
                                        """)

                                        conn_ef_db.commit()
                                        print("Table 'fyi_notices_letter' created successfully.")
                                        print()

                                    if file_exists:
                                        print("File already exists.")

                                        if file_exists:
                                            existing_data = pd.read_csv(csv_file_path)
                                            existing_users = existing_data['proceedingReqId'].tolist()

                                            for index, row in Information.iterrows():
                                                if row['proceedingReqId'] in existing_users:
                                                    print(f"FYI NOTICES LETTER: {row['panNum']} {row['proceedingReqId']} ALREADY EXISTED.")

                                                else:
                                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                                        print(f"FYI NOTICE LETTER: {row['panNum']} {row['proceedingReqId']} INSERTED SUCCESSFULLY.")
                                                    
                                                    curr_ef_db.execute("""
                                                    INSERT INTO fyi_notices_letter (
                                                        proceedingReqId, panNum, userName, loggedInUserId, noticeSection, documentRefId, 
                                                        description, issuedOn, issudedOnTime, servedOn, servedOnTime, responseDueDate, 
                                                        responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, 
                                                        responseViewedByAO, proceedingName, assessmentYear, noticeId, cc, 
                                                        mailBody, docNam, headerSeqNo, procdngReqId, applnId, date, dateTime, frommail, subject, tomail
                                                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                                    """, (
                                                        row['proceedingReqId'], row['panNum'], row['userName'], row['loggedInUserId'], row['noticeSection'], row['documentRefId'], 
                                                        row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                                        row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'],
                                                        row['responseViewedByAO'], row['proceedingName'], row['assessmentYear'], row['noticeId'], row['cc'], 
                                                        row['mailBody'], row['docNam'], row['headerSeqNo'], row['procdngReqId'], row['applnId'], row['date'], row['dateTime'], row['from'], row['subject'], row['to']
                                                    ))

                                                    conn_ef_db.commit()
                                                    print(f"FYI Notice Letter: {row['panNum']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                                    print()

                                    else:
                                        Information.to_csv(csv_file_path, header=True, index=False)
                                        print("FYI NOTICE LETTER: CSV CREATED SUCCESSFULLY.")

                                        for index, row in Information.iterrows():
                                            curr_ef_db.execute("""
                                            INSERT INTO fyi_notices_letter (
                                                proceedingReqId, panNum, userName, loggedInUserId, noticeSection, documentRefId, 
                                                description, issuedOn, issudedOnTime, servedOn, servedOnTime, responseDueDate, 
                                                responseDueDateTime, lastResponseSubmittedOn, lastResponseSubmittedOnTime, 
                                                responseViewedByAO, proceedingName, assessmentYear, noticeId, cc, 
                                                mailBody, docNam, headerSeqNo, procdngReqId, applnId, date, dateTime, frommail, subject, tomail
                                            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
                                            """, (
                                                row['proceedingReqId'], row['panNum'], row['userName'], row['loggedInUserId'], row['noticeSection'], row['documentRefId'], 
                                                row['description'], row['issuedOn'], row['issudedOnTime'], row['servedOn'], row['servedOnTime'], row['responseDueDate'],
                                                row['responseDueDateTime'], row['lastResponseSubmittedOn'], row['lastResponseSubmittedOnTime'],
                                                row['responseViewedByAO'], row['proceedingName'], row['assessmentYear'], row['noticeId'], row['cc'], 
                                                row['mailBody'], row['docNam'], row['headerSeqNo'], row['procdngReqId'], row['applnId'], row['date'], row['dateTime'], row['from'], row['subject'], row['to']
                                            ))

                                            conn_ef_db.commit()
                                            print(f"FYI Notice Letter: {row['panNum']} INSERTED SUCCESSFULLY INTO DATABASE.")
                                            print()

                                except Exception as e:
                                    print(f"No Notice letter for {response_entity}")

                except Exception as e:
                    print(f"{username} An error occurred for FYI Notice Letter : {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fyi_notice_letter
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI NOTICES LETTER: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYI NOTICES LETTER: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYI NOTICE LETTER: ERROR USERS CSV CREATED.")
                #------------------------------FYI ALL NOTICES LETTER END------------------------------

                #------------------------------FYI NOTICE DOWNLOAD START------------------------------
                try:
                    time.sleep(3)
                    if response_fyi.get('eproceedingRequests', []):
                        for proceeding_request in response_fyi.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)

                            payload_view_notices = json.dumps({
                                "serviceName": "eProceedingDetailsService",
                                "proceedingReqId": proceedingReqId,
                                "pan": response_entity,
                                "header": {
                                    "formName": "FO-041_PCDNG"
                                    }
                                })
                            
                            headers_view_notices = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            'sn': 'eProceedingDetailsService'
                            }
                            
                            response_view_notices = requests.request("POST", fyi_all_notices_url, headers=headers_view_notices, data=payload_view_notices)
                            time.sleep(3)

                            if response_view_notices.status_code==200:
                                try:
                                    view_notice = json.loads(response_view_notices.text)

                                    num_entries = len(view_notice)
                                    
                                    time.sleep(3)

                                    for entry in view_notice:
                                        headerSeqNo = entry.get("headerSeqNo", "")
                                        payload_noticeDownload = json.dumps({
                                        "serviceName": "noticeletterpdf",
                                        "headerSeqNo": headerSeqNo,
                                        "procdngReqId": proceedingReqId,
                                        "loggedInUserId": response_entity,
                                        "header": {
                                            "formName": "FO-041_PCDNG"
                                            }
                                        })
                    
                                        headers_noticeDownload = {
                                        'Content-Type': 'application/json',
                                        'Cookie': 'AuthToken='+auth_token,
                                        }
                    
                                        DownloadNotice = requests.request("POST", fyi_download_notice_url, headers = headers_noticeDownload, data = payload_noticeDownload)
                                        time.sleep(3)

                                        downloadnotice = json.loads(DownloadNotice.text)
                                        
                                        satDocId = downloadnotice.get("satDocId")
                                        doc_name = downloadnotice.get("docNam")

                                        if satDocId:
                                            pdf_url = f"https://eportal.incometax.gov.in/iec/document/{satDocId}"

                                            headers_download = {
                                            'Cookie': 'AuthToken='+auth_token,
                                            'Referer': 'https://eportal.incometax.gov.in/iec/foservices/'
                                            }

                                            response = requests.get(pdf_url, headers=headers_download)

                                            if response.status_code == 200:
                                                file_path = os.path.join(fyi_download_dir, f"{response_entity}_{doc_name}")
                                                
                                                with open(file_path, 'wb') as file:
                                                    file.write(response.content)
                                                print(f'Document download for {response_entity}:', headerSeqNo)
                                
                                except Exception as e:
                                    print(f"Direct Notice for {response_entity}")
                
                except Exception as e:
                    print(f"{username} An error occurred for FYI Notice Download : {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_fyi_notice_download
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI NOTICES DOWNLOAD: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYI NOTICES DOWNLOAD: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYI NOTICE DOWNLOAD: ERROR USERS CSV CREATED.")
                #------------------------------FYI NOTICE DOWNLOAD END------------------------------
                
                #------------------------------FYI DIRECT NOTICE DOWNLOAD START------------------------------
                try:
                    time.sleep(3)
                    if response_fyi.get('eproceedingRequests', []):
                        for proceeding_request in response_fyi.get("eproceedingRequests", []):
                            proceedingReqId = proceeding_request.get("proceedingReqId", "")

                            time.sleep(3)

                            payload_direct_download_notice = json.dumps({
                                "serviceName":"downloadClosureOrder",
                                "procdngReqId":proceedingReqId,
                                "loggedInUserId":response_entity,
                                "header":{
                                    "formName":"FO-041_PCDNG"
                                    }
                                })
                            
                            headers_payload_direct_download_notice = {
                            'Content-Type': 'application/json',
                            'Cookie': 'AuthToken='+auth_token,
                            }
                            
                            Response_Direct_Notice = requests.request("POST", fyi_direct_notice_download_url, headers=headers_payload_direct_download_notice, data=payload_direct_download_notice)
                            time.sleep(3)
                            
                            response_direct_notice= json.loads(Response_Direct_Notice.text)

                            if response_direct_notice.get("satDocDetlList"):
                                sat_doc_detl_list = response_direct_notice.get('satDocDetlList', [])
                                for entry in sat_doc_detl_list:
                                    satDocId = entry.get("satDocId")
                                    doc_name = entry.get("docNam")

                                    if satDocId:
                                        pdf_url = f"https://eportal.incometax.gov.in/iec/document/{satDocId}"

                                        headers_download = {
                                        'Cookie': 'AuthToken='+auth_token,
                                        'Referer': 'https://eportal.incometax.gov.in/iec/foservices/'
                                        }

                                        response = requests.get(pdf_url, headers=headers_download)

                                        if response.status_code == 200:
                                            file_path = os.path.join(fyi_download_dir, f"{response_entity}_{doc_name}")
                                                
                                            with open(file_path, 'wb') as file:
                                                file.write(response.content)
                                            print(f'Document download for {response_entity}:', proceedingReqId)

                except Exception as e:
                    print(f"{username} An error occurred for FYI direct Notice Download : {e}")

                    name = username
                    passwrd = password

                    user_basic_info_list = [{ 
                        "ID":name,
                        "Password":passwrd
                    }]

                    Information = pd.DataFrame(user_basic_info_list)
                    Information.replace('\x00', '', regex=True, inplace=True)

                    csv_file_path = Error_direct_notice_download_path
                    file_exists = os.path.exists(csv_file_path)
                
                    if file_exists:
                        print("File already exists.")
                                    
                        if file_exists:
                            existing_data = pd.read_csv(csv_file_path)
                            existing_users = existing_data['ID'].tolist()

                            for index, row in Information.iterrows():
                                if row['ID'] in existing_users:
                                    print(f"FYI DIRECT NOTICES DOWNLOAD: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                                else:
                                    with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                        Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                        print(f"FYI DIRECT NOTICES DOWNLOAD: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
                
                    else:
                        Information.to_csv(csv_file_path, header=True, index=False)
                        print(f"FYI DIRECT NOTICE DOWNLOAD: ERROR USERS CSV CREATED.")
                #------------------------------FYI DIRECT NOTICE DOWNLOAD END------------------------------

        except Exception as e:
            print(f"{username} An error occurred for FYI: {e}")

            name = username
            passwrd = password

            user_basic_info_list = [{ 
                "ID":name,
                "Password":passwrd
            }]

            Information = pd.DataFrame(user_basic_info_list)
            Information.replace('\x00', '', regex=True, inplace=True)

            csv_file_path = Error_fyi
            file_exists = os.path.exists(csv_file_path)
        
            if file_exists:
                print("File already exists.")
                            
                if file_exists:
                    existing_data = pd.read_csv(csv_file_path)
                    existing_users = existing_data['ID'].tolist()

                    for index, row in Information.iterrows():
                        if row['ID'] in existing_users:
                            print(f"FYI: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                        else:
                            with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                                Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                                print(f"FYI: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
            else:
                Information.to_csv(csv_file_path, header=True, index=False)
                print(f"FYI: ERROR USERS CSV CREATED.")
        #----------------------------------------FYI END------------------------------------------

    #----------------------------------------EXCEPTION BLOCK----------------------------------------        
    except Exception as e:
        print(f"{username} An error occurred for process login: {e}")

        name = username
        passwrd = password

        user_basic_info_list = [{ 
        "ID":name,
        "Password":passwrd
        }]

        Information = pd.DataFrame(user_basic_info_list)
        Information.replace('\x00', '', regex=True, inplace=True)

        csv_file_path = Error_process_login
        file_exists = os.path.exists(csv_file_path)
        
        if file_exists:
            print("File already exists.")
                            
            if file_exists:
                existing_data = pd.read_csv(csv_file_path)
                existing_users = existing_data['ID'].tolist()

                for index, row in Information.iterrows():
                    if row['ID'] in existing_users:
                        print(f"PROCESS LOGIN: ERROR USER RECORD {row['ID']} ALREADY EXISTED.")

                    else:
                        with open(csv_file_path, mode='a', newline='', encoding='utf-8') as csvfile:
                            Information.iloc[[index]].to_csv(csvfile, header=False, index=False)
                            print(f"PROCESS LOGIN: ERROR USER RECORD {row['ID']} INSERTED SUCCESSFULLY.")
        
        else:
            Information.to_csv(csv_file_path, header=True, index=False)
            print(f"PROCESS LOGIN: ERROR USERS CSV CREATED.")

    #----------------------------------------FINAL BLOCK----------------------------------------
    finally:
        curr_ef_db.close()
        conn_ef_db.close()
        
        payload_logout = json.dumps({
        "serviceName": "logoutService",
        "entity": response_entity,
        "userType": "IND"
        })

        headers_logout = {
        'Accept': 'application/json, text/plain, */*',
        'Content-Type': 'application/json',
        'Cookie': 'AuthToken='+auth_token,
        'Origin': 'https://eportal.incometax.gov.in',
        'sn': 'logoutService'
        }

        logout = requests.request("POST", logout_url, headers=headers_logout, data=payload_logout)
        time.sleep(5)
        print(f"*************************{username} User Logout Successfully*************************")

def main():
    with open(login_csv, 'r')as csvfile:
        reader = csv.reader(csvfile)
        next(reader)
        data = [(row[0], row[1]) for row in reader]

    with concurrent.futures.ProcessPoolExecutor() as executor:
        executor.map(process_login, data)

    end = time.perf_counter()
    print(f"Totla time {(end-start)} seconds")

if __name__ == "__main__":
    main()